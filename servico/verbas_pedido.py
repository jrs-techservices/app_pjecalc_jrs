# -*- coding: utf-8 -*-
"""Bloco `verbas` (e `reflexosManuais`) do .json de lançamentos — contrato da §5 do VALORES-ACEITOS-PJECALC.md.

A porta valida tudo o que dá para saber antes de abrir o PJe-Calc; o lançamento reaproveita o fluxo da planilha de
verbas: as verbas viram uma planilha temporária (TBVERBABD, TBREFLEXA, DEVIDOBD_EXP, PAGOBD_EXP) no formato que a
classe Verbas já lê em produção.
"""
import re
from pathlib import Path

from Calculo.leitura_planilha import data_da_planilha, numero_da_planilha

# Lista do Lançamento Expresso lida na tela do PJe-Calc 2.16.0 desta máquina em 07/10/2026 (54 itens). Ela muda quando o
# PJe-Calc atualiza as tabelas (em 30/09 havia 'REPOUSO SEMANAL REMUNERADO' simples): reler na tela se a porta recusar item real.
EXPRESSO = (
    "13º SALÁRIO", "ABONO PECUNIÁRIO", "ACORDO (MERA LIBERALIDADE)", "ACORDO (MULTA)",
    "ACORDO (VERBAS INDENIZATÓRIAS)", "ACORDO (VERBAS REMUNERATÓRIAS)", "ADICIONAL DE HORAS EXTRAS 50%",
    "ADICIONAL DE INSALUBRIDADE 10%", "ADICIONAL DE INSALUBRIDADE 20%", "ADICIONAL DE INSALUBRIDADE 40%",
    "ADICIONAL DE PERICULOSIDADE 30%", "ADICIONAL DE PRODUTIVIDADE 30%", "ADICIONAL DE RISCO 40%",
    "ADICIONAL DE SOBREAVISO", "ADICIONAL DE TRANSFERÊNCIA 25%", "ADICIONAL NOTURNO 20%", "AJUDA DE CUSTO",
    "AVISO PRÉVIO", "CESTA BÁSICA", "COMISSÃO", "DEVOLUÇÃO DE DESCONTOS INDEVIDOS", "DIFERENÇA SALARIAL",
    "DIÁRIAS - INTEGRAÇÃO AO SALÁRIO", "DIÁRIAS - PAGAMENTO", "FERIADO EM DOBRO", "FÉRIAS + 1/3", "GORJETA",
    "GRATIFICAÇÃO DE FUNÇÃO", "GRATIFICAÇÃO POR TEMPO DE SERVIÇO", "HORAS EXTRAS 100%", "HORAS EXTRAS 50%",
    "HORAS IN ITINERE", "INDENIZAÇÃO ADICIONAL", "INDENIZAÇÃO PIS - ABONO SALARIAL", "INDENIZAÇÃO POR DANO ESTÉTICO",
    "INDENIZAÇÃO POR DANO MATERIAL", "INDENIZAÇÃO POR DANO MORAL", "INTERVALO INTERJORNADAS", "INTERVALO INTRAJORNADA",
    "MULTA CONVENCIONAL", "MULTA DO ARTIGO 477 DA CLT", "PARTICIPAÇÃO NOS LUCROS OU RESULTADOS - PLR",
    "PRÊMIO PRODUÇÃO", "REPOUSO SEMANAL REMUNERADO (COMISSIONISTA)", "REPOUSO SEMANAL REMUNERADO EM DOBRO",
    "RESTITUIÇÃO / INDENIZAÇÃO DE DESPESA", "SALDO DE EMPREITADA", "SALDO DE SALÁRIO", "SALÁRIO MATERNIDADE",
    "SALÁRIO RETIDO", "TÍQUETE-ALIMENTAÇÃO", "VALE TRANSPORTE", "VALOR PAGO - NÃO TRIBUTÁVEL",
    "VALOR PAGO - TRIBUTÁVEL",
)

# código do XML do .PJC → código da planilha de verbas (o que a classe Verbas entende)
PARCELA = {"FIXA": "F", "VARIAVEL": "V"}
VALOR = {"CALCULADO": "C", "INFORMADO": "I"}
CARACTERISTICA = {"COMUM": "C", "DECIMO_TERCEIRO_SALARIO": "DT", "AVISO_PREVIO": "AP", "FERIAS": "F"}
OCORRENCIA = {"DESLIGAMENTO": "DL", "DEZEMBRO": "DZ", "MENSAL": "M", "PERIODO_AQUISITIVO": "PA"}
OCORRENCIA_DA_CARACTERISTICA = {"COMUM": "MENSAL", "DECIMO_TERCEIRO_SALARIO": "DEZEMBRO", "AVISO_PREVIO": "DESLIGAMENTO",
                                "FERIAS": "PERIODO_AQUISITIVO"}
GERA = {"DEVIDO": "DV", "DIFERENCA": "DF"}
BASE_TABELADA = {"HISTORICO_SALARIAL": "HS", "MAIOR_REMUNERACAO": "MR", "PISO_SALARIAL": "PS", "SALARIO_MINIMO": "SM",
                 "VALE_TRANSPORTE": "VT"}
DIVISOR = {"OUTRO_VALOR": "IN", "CARGA_HORARIA": "CH", "DIAS_UTEIS": "DU", "IMPORTADA_DO_CARTAO": "ICP"}
# AVOS: na tela de 13º/Férias a 2ª opção é "Avos" (código AV da planilha); APURADA: o PJe-Calc já marca no Aviso Prévio.
QUANTIDADE = {"INFORMADA": "IN", "IMPORTADA_DO_CALENDARIO": "ICL", "IMPORTADA_DO_CARTAO": "ICP", "AVOS": "AV", "APURADA": None}
CALENDARIO = {"REPOUSOS_E_FERIADOS": "RF", "REPOUSOS": "R", "FERIADOS": "F", "DIAS_UTEIS": "DU"}
VALOR_PAGO = {"INFORMADO": "I", "CALCULADO": "C"}
COMPORTAMENTO = {"VALOR_MENSAL": "VM", "MEDIA_PELO_VALOR": "MA", "MEDIA_PELO_VALOR_CORRIGIDO": "MC", "MEDIA_PELA_QUANTIDADE": "MQ"}
PERIODO_MEDIA = {"PERIODO_AQUISITIVO": "PA", "ANO_CIVIL": "AC", "ULTIMOS_DOZE_MESES_DO_CONTRATO": "DM",
                 "DOZE_MESES_ANTERIORES_AO_VENCIMENTO_DA_PARCELA": "DV"}
PERIODOS_DA_OCORRENCIA = {
    "MENSAL": (), "DEZEMBRO": ("ANO_CIVIL",),
    "DESLIGAMENTO": ("ULTIMOS_DOZE_MESES_DO_CONTRATO", "DOZE_MESES_ANTERIORES_AO_VENCIMENTO_DA_PARCELA"),
    "PERIODO_AQUISITIVO": ("PERIODO_AQUISITIVO", "DOZE_MESES_ANTERIORES_AO_VENCIMENTO_DA_PARCELA"),
}
FRACAO = {"MANTER": "M", "INTEGRALIZAR": "I", "DESPREZAR": "D", "DESPREZAR_MENOR_QUE_15_DIAS": "DMQ"}
REFLEXOS = {
    "decimoTerceiro": "SFLREFLEXO13TERCEIRO", "ferias": "SFLREFLEXOFERIAS", "aviso": "SFLREFLEXOAVISO",
    "repouso": "SFLREFLEXOREPOUSO", "art477": "SFLREFLEXOART477CLT", "art467": "SFLREFLEXOART467CLT",
    "separarFerias": "SFLSEPARARFERIAS", "horasExtras50": "SFLREFLEXOHEXTRAS50", "horasExtras100": "SFLREFLEXOHEXTRAS100",
    "gratificacaoAnual": "SFLREFLEXOGANUAL", "gratificacaoSemestral": "SFLREFLEXOGSEMESTRAL",
    "licencaPremio": "SFLREFLEXOLPREMIO", "apip": "SFLREFLEXOAPIP", "plr": "SFLREFLEXOPLR",
}
# o PJe-Calc sempre oferece estes na lista de reflexos da principal; os demais precisam de reflexo manual
SEMPRE_OFERECIDOS = {"decimoTerceiro", "ferias", "aviso"}
INCIDENCIAS = {"inss": "SFLINCIDENCIAINSS", "irpf": "SFLINCIDENCIAIRPF", "fgts": "SFLINCIDENCIAFGTS",
               "previdenciaPrivada": "SFLINCIDENCIAPREVPRIVADA", "pensaoAlimenticia": "SFLINCIDENCIAPENSAO"}
MARCACOES = {"jurosSumula439": "STPJUROSSUMULA", "comporPrincipal": "SFLCOMPORPRINCIPAL",
             "zerarValorNegativo": "SFLPADRAOZERARVALORNEGATIVO", "dobrarValor": "SFLDOBRARVALOR",
             "excluirFaltaJustificada": "SFLEXCLUIRFALTAJUSTIFICADA", "excluirFaltaNaoJustificada": "SFLEXCLUIRFALTANAOJUSTIFICADA",
             "excluirFeriasGozadas": "SFLEXCLUIRFERIASGOZADAS"}
ASSUNTO_CNJ_PADRAO = 2581  # o assunto de todos os reflexos manuais da planilha de verbas v1.12

# colunas na ordem da planilha de verbas v1.11/v1.12: a classe Verbas preenche a tela nesta ordem
COLUNAS_VERBA = [
    "IIDVERBAJRS", "SNMVERBAJRS", "SNMVERBAEXPRESSOPJECALC", "SNMDESCRICAOVERBA", "SFLINCIDENCIAFGTS", "SFLINCIDENCIAIRPF",
    "SFLINCIDENCIAINSS", "SFLREFLEXO13TERCEIRO", "SFLREFLEXOFERIAS", "SFLREFLEXOAVISO", "SFLREFLEXOREPOUSO",
    "SFLREFLEXOART477CLT", "SFLREFLEXOART467CLT", "SFLINCIDENCIAPREVPRIVADA", "SFLINCIDENCIAPENSAO", "STPCARACTERISTICAVERBA",
    "STPOCORRENCIAPAGAMENTO", "ICDASSUNTO", "STPPARCELA", "STPVALOR", "STPJUROSSUMULA", "STPTIPOVERBA", "STPGERARVERBAREFLEXA",
    "STPGERARVERBAPRINCIPAL", "SFLCOMPORPRINCIPAL", "SFLPADRAOZERARVALORNEGATIVO", "DDTINICIO", "DDTTERMINO",
    "SFLEXCLUIRFALTAJUSTIFICADA", "SFLEXCLUIRFALTANAOJUSTIFICADA", "SFLEXCLUIRFERIASGOZADAS", "SFLDOBRARVALOR", "RVLVALORDEVIDO",
    "SFLPROPORCIONALIDADE", "STPBASEDEVIDO", "STPDIVISOR", "RVLOUTRODIVISOR", "STPCARTAOPONTODIVISOR", "RVLMULTIPLICADOR", "STPQUANTIDADE",
    "STPCALENDARIO", "STPCARTAOPONTO", "RVLOUTROVALORQTD", "SFLQTDPROPORCIONALIDADE", "STPVALORPAGO", "STPBASEPAGO",
    "RVLINOUTROVALOR", "SFLVPGPROPORCIONALIDADE", "RVLQUANTIDADE", "SDSCOMENTARIO", "STPCOMPORTAMENTOREFLEXO",
    "STPPERIODOMEDIAREFLEXO", "STPTRATAMENTOFRACAOMESREFLEXO", "SFLSEPARARFERIAS", "SFLREFLEXOHEXTRAS50",
    "SFLREFLEXOHEXTRAS100", "SFLREFLEXOGANUAL", "SFLREFLEXOGSEMESTRAL", "SFLREFLEXOLPREMIO", "SFLREFLEXOAPIP", "SFLREFLEXOPLR",
    "SFLREFLEXOSFGTS",  # só nesta planilha: FGTS nos reflexos ativos (menos férias), do reflexosFgts do .json
]
COLUNAS_REFLEXA = [
    "IIDVERBAREFLEXA", "SNMVERBAREFLEXA", "SNMVERBAEXPRESSOPJECALC", "SNMDESCRICAOVERBAREFLEXA", "ICDASSUNTO", "STPPARCELA",
    "STPVALOR", "SFLINCIDENCIAINSS", "SFLINCIDENCIAIRPF", "SFLINCIDENCIAFGTS", "SFLINCIDENCIAPREVPRIVADA", "SFLINCIDENCIAPENSAO",
    "STPCARACTERISTICAVERBA", "STPOCORRENCIAPAGAMENTO", "STPJUROSSUMULA", "STPTIPOVERBA", "STPGERARVERBAREFLEXA",
    "STPGERARVERBAPRINCIPAL", "SFLCOMPORPRINCIPAL", "SFLPADRAOZERARVALORNEGATIVO", "DDTINICIO", "DDTTERMINO",
    "SFLEXCLUIRFALTAJUSTIFICADA", "SFLEXCLUIRFALTANAOJUSTIFICADA", "SFLEXCLUIRFERIASGOZADAS", "SFLDOBRARVALOR", "RVLVALORDEVIDO",
    "SFLPROPORCIONALIDADE", "STPBASEDEVIDO", "STPDIVISOR", "RVLOUTRODIVISOR", "STPCARTAOPONTODIVISOR", "RVLMULTIPLICADOR", "STPQUANTIDADE",
    "STPCALENDARIO", "STPCARTAOPONTO", "RVLOUTROVALORQTD", "SFLQTDPROPORCIONALIDADE", "STPVALORPAGO", "STPBASEPAGO",
    "RVLINOUTROVALOR", "SFLVPGPROPORCIONALIDADE", "RVLQUANTIDADE", "SDSCOMENTARIO", "SFLREFLEXO13TERCEIRO", "SFLREFLEXOFERIAS",
    "SFLREFLEXOAVISO", "SFLREFLEXOREPOUSO", "SFLREFLEXOART477CLT", "SFLREFLEXOART467CLT", "STPCOMPORTAMENTOREFLEXO",
    "STPPERIODOMEDIAREFLEXO", "STPTRATAMENTOFRACAOMESREFLEXO",
    "FKVERBAPRINCIPAL",  # só nesta planilha: amarra o reflexo manual à principal (as planilhas da casa não têm)
]


# O PJe-Calc grava texto em Latin-1: o que está fora dele vira "¿" (ex.: o travessão "—").
SUBSTITUTOS = {"—": "-", "–": "-", "‒": "-", "−": "-", "“": '"', "”": '"', "„": '"',
               "‘": "'", "’": "'", "´": "'", "…": "...", " ": " ", "•": "-"}
LIMITE_DESCRICAO = 50  # o PJe-Calc corta a descrição da verba em 50 caracteres (visto nos .PJC da auditoria do SAP)


def texto_latin1(texto):
    """Troca o que o PJe-Calc não guarda por equivalente simples; o que não tiver equivalente é retirado."""
    texto = "".join(SUBSTITUTOS.get(c, c) for c in str(texto))
    return "".join(c for c in texto if ord(c) < 256)


def normalizar_descricao(texto):
    return " ".join(texto_latin1(texto).split()).upper()


def _descricao(r, ler, campo="descricao"):
    original = " ".join(str(ler.item.get(campo) or "").split()).upper()
    descricao = normalizar_descricao(original)
    if descricao != original:
        r.avisos.append(f"{ler.base}.{campo}: caracteres que o PJe-Calc não grava foram trocados: {original!r} -> {descricao!r}")
    return descricao


class _Leitor:
    """Lê um item (verba ou reflexo manual) acumulando erros em `r` com o caminho do campo."""

    def __init__(self, r, base, item):
        self.r, self.base, self.item = r, base, item

    def erro(self, campo, mensagem, aceitos=None):
        self.r.erro(f"{self.base}.{campo}", self._valor(campo), mensagem, aceitos)

    def _valor(self, campo):
        atual = self.item
        for parte in campo.split("."):
            atual = atual.get(parte) if isinstance(atual, dict) else None
        return atual

    def lista(self, campo, tabela, obrigatorio=False):
        valor = self._valor(campo)
        if valor in (None, ""):
            if obrigatorio:
                self.erro(campo, "obrigatório", sorted(tabela))
            return None
        if valor not in tabela:
            self.erro(campo, "valor não reconhecido", sorted(tabela))
            return None
        return valor

    def booleano(self, campo):
        valor = self._valor(campo)
        if valor is None:
            return None
        if not isinstance(valor, bool):
            self.erro(campo, "use true ou false")
            return None
        return valor

    def numero(self, campo, obrigatorio=False):
        valor = self._valor(campo)
        if valor in (None, ""):
            if obrigatorio:
                self.erro(campo, "obrigatório")
            return None
        numero = numero_da_planilha(valor)
        if numero is None:
            self.erro(campo, "deve ser número")
        return numero

    def data(self, campo):
        valor = self._valor(campo)
        if valor in (None, ""):
            return None
        data = data_da_planilha(valor)
        if data is None:
            self.erro(campo, "data inválida: use AAAA-MM-DD")
            return None
        return data.strftime("%d/%m/%Y")


def _sn(valor):
    return None if valor is None else ("S" if valor else "N")


def _parametros_comuns(ler, linha):
    """Campos que verba principal e reflexo manual têm em comum; preenche `linha` com os códigos da planilha."""
    parcela = ler.lista("parcela", PARCELA)
    valor = ler.lista("valor", VALOR, obrigatorio=True)
    caracteristica = ler.lista("caracteristica", CARACTERISTICA)
    ocorrencia = ler.lista("ocorrenciaPagamento", OCORRENCIA)
    linha.update({"STPPARCELA": PARCELA.get(parcela), "STPVALOR": VALOR.get(valor),
                  "STPCARACTERISTICAVERBA": CARACTERISTICA.get(caracteristica), "STPOCORRENCIAPAGAMENTO": OCORRENCIA.get(ocorrencia)})
    incidencias = ler.item.get("incidencias")
    if incidencias is not None and not isinstance(incidencias, dict):
        ler.erro("incidencias", "deve ser um objeto { inss, irpf, fgts, previdenciaPrivada, pensaoAlimenticia }")
    elif incidencias:
        for chave in incidencias:
            if chave not in INCIDENCIAS:
                ler.erro(f"incidencias.{chave}", "incidência não reconhecida", sorted(INCIDENCIAS))
        for chave, coluna in INCIDENCIAS.items():
            linha[coluna] = _sn(ler.booleano(f"incidencias.{chave}"))
    for chave, coluna in MARCACOES.items():
        linha[coluna] = _sn(ler.booleano(chave))
    linha["STPGERARVERBAREFLEXA"] = GERA.get(ler.lista("geraReflexo", GERA))
    linha["STPGERARVERBAPRINCIPAL"] = GERA.get(ler.lista("gerarPrincipal", GERA))
    periodo = ler.item.get("periodo")
    if periodo is not None and not isinstance(periodo, dict):
        ler.erro("periodo", "deve ser { inicio, termino } em AAAA-MM-DD")
    linha["DDTINICIO"] = ler.data("periodo.inicio")
    linha["DDTTERMINO"] = ler.data("periodo.termino")
    _conferir_prescricao(ler, linha["DDTINICIO"])
    if valor == "INFORMADO":
        linha["RVLVALORDEVIDO"] = ler.numero("valorDevido", obrigatorio=True)
        linha["SFLPROPORCIONALIDADE"] = _sn(ler.booleano("proporcionalidade"))
    elif valor == "CALCULADO":
        linha["SFLPROPORCIONALIDADE"] = _sn(ler.booleano("proporcionalidade"))
        _divisor_quantidade(ler, linha, caracteristica)
    _valor_pago(ler, linha)
    comentario = ler.item.get("comentario")
    if comentario:
        linha["SDSCOMENTARIO"] = str(comentario)
    return valor, caracteristica, ocorrencia


COLUNAS_DO_CARTAO_PADRAO = ("Dias Trabalhados", "Hs EXT", "Hs Trabalhadas")  # o que a apuração com os critérios padrão gera


def _coluna_do_cartao(ler, campo, tipo):
    """divisor/quantidade IMPORTADA_DO_CARTAO: `coluna` (texto do seletor da tela) ou `colunas` (várias, que o PJe-Calc
    soma). Vai para a planilha separado por '|'."""
    coluna, colunas = ler._valor(f"{campo}.coluna"), ler._valor(f"{campo}.colunas")
    if tipo != "IMPORTADA_DO_CARTAO":
        for nome, valor in (("coluna", coluna), ("colunas", colunas)):
            if valor not in (None, "", []):
                ler.erro(f"{campo}.{nome}", "só com tipo IMPORTADA_DO_CARTAO")
        return None
    if coluna not in (None, "") and colunas not in (None, []):
        ler.erro(f"{campo}.colunas", "use `coluna` ou `colunas`, não os dois")
        return None
    lista = [coluna] if coluna not in (None, "") else colunas
    if not isinstance(lista, list) or not lista or not all(isinstance(c, str) and c.strip() and "|" not in c for c in lista):
        ler.erro(f"{campo}.coluna", "obrigatória com tipo IMPORTADA_DO_CARTAO: a coluna da apuração do cartão "
                                    "(ou `colunas`, uma lista de textos)", COLUNAS_DO_CARTAO_PADRAO)
        return None
    return "|".join(" ".join(c.split()) for c in lista)


# Expressos que nascem com ocorrência Desligamento: num cálculo sem data de demissão o PJe-Calc 2.16 recusa criá-los
# ("Para a ocorrência de desligamento, a data de demissão deve ser informada nos parâmetros do cálculo"); sonda de
# 29/10/2026 com os 54 Expressos (17583).
EXPRESSOS_DE_DESLIGAMENTO = {
    "VALOR PAGO - NÃO TRIBUTÁVEL", "VALOR PAGO - TRIBUTÁVEL", "INDENIZAÇÃO ADICIONAL", "INDENIZAÇÃO PIS - ABONO SALARIAL",
    "MULTA CONVENCIONAL", "MULTA DO ARTIGO 477 DA CLT", "AVISO PRÉVIO", "RESTITUIÇÃO / INDENIZAÇÃO DE DESPESA",
    "SALDO DE EMPREITADA", "SALDO DE SALÁRIO",
}


def _conferir_prescricao(ler, inicio):
    """Com a prescrição quinquenal, o PJe-Calc recusa a verba que começa antes de 5 anos do ajuizamento: "A data de início
    informada não pode ser anterior à data da prescrição quinquenal (5 anos antes da Data do Ajuizamento)" (17584, 28/10/2026)."""
    from datetime import datetime, timedelta
    ajuizamento = ler.r.valores.get("data_ajuizamento")
    if not inicio or ler.r.valores.get("aplicar_prescricao_verbas") != "True" or not isinstance(ajuizamento, int):
        return
    data = datetime(1899, 12, 30) + timedelta(days=ajuizamento)
    try:
        limite = data.replace(year=data.year - 5)
    except ValueError:  # 29/02
        limite = data.replace(year=data.year - 5, day=28)
    if datetime.strptime(inicio, "%d/%m/%Y") < limite:
        ler.erro("periodo.inicio", f"anterior à prescrição quinquenal ({limite:%d/%m/%Y}, 5 anos antes do ajuizamento): "
                                   f"o PJe-Calc recusa a verba; use {limite:%Y-%m-%d} ou depois")


def _divisor_quantidade(ler, linha, caracteristica):
    # Divisor e multiplicador são obrigatórios na tela de verba CALCULADO: sem eles o PJe-Calc não salva a verba.
    divisor = ler.item.get("divisor")
    if divisor is None:
        ler.erro("divisor", "obrigatório em verba CALCULADO ({ tipo, valor })", sorted(DIVISOR))
    else:
        if not isinstance(divisor, dict):
            ler.erro("divisor", "deve ser { tipo, valor }")
        else:
            tipo = ler.lista("divisor.tipo", DIVISOR, obrigatorio=True)
            linha["STPDIVISOR"] = DIVISOR.get(tipo)
            linha["RVLOUTRODIVISOR"] = ler.numero("divisor.valor", obrigatorio=tipo == "OUTRO_VALOR")
            linha["STPCARTAOPONTODIVISOR"] = _coluna_do_cartao(ler, "divisor", tipo)
    linha["RVLMULTIPLICADOR"] = ler.numero("multiplicador", obrigatorio=True)
    quantidade = ler.item.get("quantidade")
    if quantidade is None:
        return
    if not isinstance(quantidade, dict):
        ler.erro("quantidade", "deve ser { tipo, valor, calendario, proporcionalidade }")
        return
    tipo = ler.lista("quantidade.tipo", QUANTIDADE, obrigatorio=True)
    if tipo == "AVOS" and caracteristica not in ("DECIMO_TERCEIRO_SALARIO", "FERIAS"):
        ler.erro("quantidade.tipo", "AVOS só com característica DECIMO_TERCEIRO_SALARIO ou FERIAS")
    if tipo == "APURADA" and caracteristica != "AVISO_PREVIO":
        ler.erro("quantidade.tipo", "APURADA só com característica AVISO_PREVIO")
    linha["STPQUANTIDADE"] = QUANTIDADE.get(tipo)
    linha["STPCARTAOPONTO"] = _coluna_do_cartao(ler, "quantidade", tipo)
    if tipo == "INFORMADA":
        linha["RVLOUTROVALORQTD"] = ler.numero("quantidade.valor", obrigatorio=True)
    calendario = ler.lista("quantidade.calendario", CALENDARIO)
    if calendario and tipo != "IMPORTADA_DO_CALENDARIO":
        ler.erro("quantidade.calendario", "só com quantidade IMPORTADA_DO_CALENDARIO")
    linha["STPCALENDARIO"] = CALENDARIO.get(calendario)
    linha["SFLQTDPROPORCIONALIDADE"] = _sn(ler.booleano("quantidade.proporcionalidade"))


def _valor_pago(ler, linha):
    pago = ler.item.get("valorPago")
    if pago is None:
        return []
    if not isinstance(pago, dict):
        ler.erro("valorPago", "deve ser { tipo, base, valor, quantidade, proporcionalidade, bases }")
        return []
    tipo = ler.lista("valorPago.tipo", VALOR_PAGO, obrigatorio=True)
    linha["STPVALORPAGO"] = VALOR_PAGO.get(tipo)
    base = ler.lista("valorPago.base", BASE_TABELADA)
    if tipo == "CALCULADO" and base is None and ler.item["valorPago"].get("base") in (None, ""):
        # o PJe-Calc recusa salvar a verba com valor pago calculado sem base (falha do 17111)
        ler.erro("valorPago.base", "obrigatória com valorPago.tipo CALCULADO", sorted(BASE_TABELADA))
    linha["STPBASEPAGO"] = BASE_TABELADA.get(base)
    linha["RVLINOUTROVALOR"] = ler.numero("valorPago.valor", obrigatorio=tipo == "INFORMADO")
    linha["RVLQUANTIDADE"] = ler.numero("valorPago.quantidade")
    linha["SFLVPGPROPORCIONALIDADE"] = _sn(ler.booleano("valorPago.proporcionalidade"))
    bases = pago.get("bases") or []
    if base == "HISTORICO_SALARIAL" and not bases:
        ler.erro("valorPago.bases", "com base HISTORICO_SALARIAL, informe os históricos em valorPago.bases[]")
    if bases and base != "HISTORICO_SALARIAL":
        ler.erro("valorPago.bases", "só com valorPago.base HISTORICO_SALARIAL")
    resultado = []
    for i, b in enumerate(bases if isinstance(bases, list) else []):
        if not isinstance(b, dict) or not str(b.get("referencia") or "").strip():
            ler.erro(f"valorPago.bases", f"item {i}: informe {{ referencia (nome do histórico), proporcionalizar }}")
            continue
        resultado.append((str(b["referencia"]).strip(), _sn(bool(b.get("proporcionalizar")))))
    ler.item["_bases_pago"] = resultado
    return resultado


def _reflexos(ler, linha):
    reflexos = ler.item.get("reflexos")
    marcados = set()
    if reflexos is None:
        return marcados
    if not isinstance(reflexos, dict):
        ler.erro("reflexos", "deve ser um objeto de booleanos", sorted(REFLEXOS))
        return marcados
    for chave, valor in reflexos.items():
        if chave not in REFLEXOS:
            ler.erro(f"reflexos.{chave}", "reflexo não reconhecido", sorted(REFLEXOS))
            continue
        if not isinstance(valor, bool):
            ler.erro(f"reflexos.{chave}", "use true ou false")
            continue
        linha[REFLEXOS[chave]] = "S" if valor else "N"
        if valor:
            marcados.add(chave)
    return marcados


def interpretar_verbas(r):
    """Valida `verbas` e `reflexosManuais`. Guarda em r.verbas_pedido o que precisa ser lançado (ou None)."""
    dados = r.dados
    r.verbas_pedido = None
    r.historicos_referenciados = []
    if "verbas" not in dados:
        if "reflexosManuais" in dados:
            r.ignorar("reflexosManuais", "sem o bloco 'verbas'")
        return
    verbas = dados["verbas"]
    if not isinstance(verbas, list):
        r.erro("verbas", type(verbas).__name__, "deve ser uma lista de verbas")
        return
    if not verbas:
        r.usados.add("verbas")
        r.avisos.append("verbas: [] — nenhuma verba a lançar; fica só a verba-modelo que o PJe-Calc precisa para liquidar.")
        return
    novas = [isinstance(v, dict) and "id" in v and "expresso" in v for v in verbas]
    if not any(novas):
        # formato antigo do emissor (tipo/ocorrencia/modo…): sem id e expresso não é o contrato; ignora como antes
        r.ignorar("verbas", "formato anterior ao contrato (itens sem 'id' e 'expresso'): bloco ignorado")
        if "reflexosManuais" in dados:
            r.ignorar("reflexosManuais", "o bloco 'verbas' foi ignorado")
        return

    erros_antes = len(r.erros)
    por_id, descricoes, refs_verba, bases_hs, marcados_por_id = {}, {}, {}, {}, {}
    for n, item in enumerate(verbas):
        base = f"verbas[{n}]"
        if not isinstance(item, dict) or "id" not in item or "expresso" not in item:
            r.erro(base, item if not isinstance(item, dict) else None, "cada verba precisa de 'id' e 'expresso'")
            continue
        ler = _Leitor(r, base, item)
        ident = str(item["id"]).strip()
        if not ident:
            ler.erro("id", "obrigatório")
            continue
        if ident in por_id:
            ler.erro("id", f"repetido: já usado em outra verba do pedido")
            continue
        if item["expresso"] not in EXPRESSO:
            ler.erro("expresso", "não está na lista do Lançamento Expresso do PJe-Calc", EXPRESSO)
        elif item["expresso"] in EXPRESSOS_DE_DESLIGAMENTO and not r.valores.get("data_rescisao"):
            ler.erro("expresso", "este Expresso nasce com ocorrência Desligamento e o PJe-Calc recusa criá-lo sem data de "
                                 "demissão (dadosDoCalculo.dataDemissao): informe a demissão ou use outro Expresso",
                     sorted(EXPRESSOS_DE_DESLIGAMENTO))
        descricao = _descricao(r, ler)
        if not descricao:
            ler.erro("descricao", "obrigatória")
        linha = {"SNMVERBAJRS": descricao, "SNMVERBAEXPRESSOPJECALC": item["expresso"], "SNMDESCRICAOVERBA": descricao,
                 "STPTIPOVERBA": "P"}
        valor, caracteristica, _ = _parametros_comuns(ler, linha)
        marcados_por_id[ident] = _reflexos(ler, linha)
        # reflexosFgts: liga a incidência FGTS em cada reflexo ativo da principal, menos férias + 1/3 (OJ 195)
        fgts_reflexos = ler.booleano("reflexosFgts")
        linha["SFLREFLEXOSFGTS"] = _sn(fgts_reflexos)
        if fgts_reflexos and not marcados_por_id[ident]:
            r.avisos.append(f"{base}.reflexosFgts: a verba não marca nenhum reflexo; não há onde ligar o FGTS")
        if "separarFerias" in marcados_por_id[ident] and "ferias" not in marcados_por_id[ident]:
            ler.erro("reflexos.separarFerias", "separarFerias precisa de reflexos.ferias = true (separa o terço do reflexo de férias)")

        bases = item.get("bases") or []
        if not isinstance(bases, list):
            ler.erro("bases", "deve ser uma lista de { tipo, referencia, proporcionalizar }")
            bases = []
        tem_hs = False
        refs_verba[ident], bases_hs[ident] = [], []
        for i, b in enumerate(bases):
            if not isinstance(b, dict) or b.get("tipo") not in ("HISTORICO", "VERBA") or not str(b.get("referencia") or "").strip():
                r.erro(f"{base}.bases[{i}]", b, "use { tipo: HISTORICO | VERBA, referencia, proporcionalizar }")
                continue
            prop = _sn(bool(b.get("proporcionalizar")))
            if b["tipo"] == "HISTORICO":
                tem_hs = True
                bases_hs[ident].append((str(b["referencia"]).strip(), prop, f"{base}.bases[{i}].referencia"))
            else:
                refs_verba[ident].append((str(b["referencia"]).strip(), prop, f"{base}.bases[{i}].referencia"))
        if bases and valor == "INFORMADO":
            ler.erro("bases", "bases só se aplicam a verba CALCULADO")
        tabelada = ler.lista("baseTabelada", BASE_TABELADA)
        if valor == "CALCULADO":
            if tabelada is None and tem_hs:
                tabelada = "HISTORICO_SALARIAL"
                r.avisos.append(f"{base}: baseTabelada derivada como HISTORICO_SALARIAL (há base de histórico).")
            elif tabelada == "HISTORICO_SALARIAL" and not tem_hs:
                ler.erro("baseTabelada", "HISTORICO_SALARIAL exige ao menos um bases[] do tipo HISTORICO")
            elif tabelada not in (None, "HISTORICO_SALARIAL") and tem_hs:
                ler.erro("baseTabelada", f"{tabelada} não convive com base do tipo HISTORICO")
            if tabelada is None and refs_verba[ident]:
                # a tela só inclui as verbas da base depois de escolher a base tabelada; sem ela, a classe Verbas não chega lá
                ler.erro("baseTabelada", "obrigatória quando há bases[] do tipo VERBA", sorted(BASE_TABELADA))
            linha["STPBASEDEVIDO"] = BASE_TABELADA.get(tabelada)

        por_id[ident] = {"linha": linha, "descricao": descricao, "item": item, "base": base}
        descricoes.setdefault(descricao, []).append(ident)

    # referências entre verbas: existem, não apontam para si, sem ciclo
    for ident, refs in refs_verba.items():
        for ref, _, caminho in refs:
            if ref == ident:
                r.erro(caminho, ref, "a verba não pode ser base de si mesma")
            elif ref not in por_id:
                r.erro(caminho, ref, "não há verba com este id no pedido")
    ordem = _ordem_sem_ciclo(r, {i: [ref for ref, _, _ in refs if ref in por_id and ref != i] for i, refs in refs_verba.items() if i in por_id})

    # nomes de histórico usados nas bases; conferidos por conferir_historicos(), que sabe de onde vêm os históricos
    r.historicos_referenciados = [(nome, caminho) for hs in bases_hs.values() for nome, _, caminho in hs]
    r.historicos_referenciados += [(nome, f"{d['base']}.valorPago.bases") for d in por_id.values() for nome, _ in d["item"].get("_bases_pago", [])]

    manuais = _reflexos_manuais(r, por_id, marcados_por_id)

    # descrição única para verba endereçada (alvo de base VERBA, com reflexo marcado ou principal de reflexo manual)
    enderecadas = {ref for refs in refs_verba.values() for ref, _, _ in refs if ref in por_id}
    enderecadas |= {i for i, m in marcados_por_id.items() if m and i in por_id}
    enderecadas |= {m["principal"] for m in manuais}
    # descrição acima do limite: o PJe-Calc corta. Para verba endereçada isso quebra a busca dos reflexos e das bases.
    for ident, dados_verba in por_id.items():
        tamanho = len(dados_verba["descricao"])
        if tamanho <= LIMITE_DESCRICAO:
            continue
        if ident in enderecadas:
            r.erro(f"{dados_verba['base']}.descricao", dados_verba["descricao"],
                   f"o PJe-Calc guarda até {LIMITE_DESCRICAO} caracteres (tem {tamanho}) e esta verba tem reflexo ou é base "
                   "de outra: cortada, a automação não a acharia")
        else:
            r.avisos.append(f"{dados_verba['base']}.descricao: tem {tamanho} caracteres; o PJe-Calc guarda os {LIMITE_DESCRICAO} primeiros "
                            f"({dados_verba['descricao'][:LIMITE_DESCRICAO]!r})")
    for ident in sorted(enderecadas):
        iguais = descricoes.get(por_id[ident]["descricao"], [])
        if len(iguais) > 1:
            r.erro(f"{por_id[ident]['base']}.descricao", por_id[ident]["descricao"],
                   f"verba referenciada precisa de descrição única no pedido; repetida nas verbas {', '.join(iguais)}")

    usa_cartao = [d["base"] for d in por_id.values() if "ICP" in (d["linha"].get("STPDIVISOR"), d["linha"].get("STPQUANTIDADE"))]
    if usa_cartao and "cartao_ponto_arquivo" not in r.valores and not any(e["campo"] == "arquivos.cartao-de-ponto" for e in r.erros):
        r.erro("arquivos.cartao-de-ponto", None, f"obrigatório: {', '.join(usa_cartao)} usa(m) divisor ou quantidade "
                                                 "IMPORTADA_DO_CARTAO")

    if len(r.erros) > erros_antes:
        return
    numero = {ident: n for n, ident in enumerate(ordem, start=1)}
    for ident in ordem:
        dados_verba = por_id[ident]
        dados_verba["linha"]["IIDVERBAJRS"] = numero[ident]
        dados_verba["devido"] = [("HS", nome, prop) for nome, prop, _ in bases_hs.get(ident, [])]
        dados_verba["devido"] += [("VB", por_id[ref]["descricao"], prop) for ref, prop, _ in refs_verba.get(ident, [])]
        dados_verba["pago"] = dados_verba["item"].get("_bases_pago", [])
    for m in manuais:
        m["linha"]["SNMVERBAEXPRESSOPJECALC"] = por_id[m["principal"]]["item"]["expresso"]
        m["linha"]["FKVERBAPRINCIPAL"] = numero[m["principal"]]
    r.usados.add("verbas")
    if "reflexosManuais" in dados:
        r.usados.add("reflexosManuais")
    r.verbas_pedido = {"verbas": [por_id[i] for i in ordem], "reflexos": manuais}
    for dados_verba in por_id.values():
        dados_verba["item"].pop("_bases_pago", None)


def conferir_historicos(r, nomes):
    """Base de histórico cujo nome não existe no pedido: erro (o PJe-Calc zeraria a verba em silêncio).

    `nomes`: os históricos que o cálculo vai ter (da planilha base, ou dos .csv do .json); None = não dá para saber.
    """
    if nomes is None or not getattr(r, "verbas_pedido", None):
        return
    existentes = {n.strip().upper() for n in nomes}
    for nome, caminho in r.historicos_referenciados:
        if nome.upper() not in existentes:
            r.erro(caminho, nome, "não há histórico salarial com este nome no pedido", sorted(nomes))


def _ordem_sem_ciclo(r, dependencias):
    """Ordem de criação: a verba usada como base nasce antes de quem a usa. Ciclo → erro."""
    ordem, estado = [], {}

    def visitar(ident, caminho):
        if estado.get(ident) == "feito":
            return
        if estado.get(ident) == "visitando":
            ciclo = caminho[caminho.index(ident):] + [ident]
            r.erro("verbas", " → ".join(ciclo), "ciclo nas bases do tipo VERBA")
            return
        estado[ident] = "visitando"
        for dep in dependencias.get(ident, []):
            visitar(dep, caminho + [ident])
        estado[ident] = "feito"
        ordem.append(ident)

    for ident in dependencias:
        visitar(ident, [])
    return ordem


def _reflexos_manuais(r, por_id, marcados_por_id):
    itens = r.dados.get("reflexosManuais") or []
    if not isinstance(itens, list):
        r.erro("reflexosManuais", type(itens).__name__, "deve ser uma lista")
        return []
    manuais, vistos = [], set()
    for n, item in enumerate(itens):
        base = f"reflexosManuais[{n}]"
        if not isinstance(item, dict):
            r.erro(base, item, "deve ser um objeto")
            continue
        ler = _Leitor(r, base, item)
        principal = str(item.get("principal") or "").strip()
        chave = item.get("reflexo")
        if principal not in por_id:
            ler.erro("principal", "não há verba com este id no pedido")
            continue
        if chave not in REFLEXOS or chave in SEMPRE_OFERECIDOS:
            ler.erro("reflexo", "reflexo sem criação manual", sorted(set(REFLEXOS) - SEMPRE_OFERECIDOS))
            continue
        if (principal, chave) in vistos:
            ler.erro("reflexo", f"já há reflexo manual '{chave}' para a verba {principal}")
            continue
        vistos.add((principal, chave))
        if chave not in marcados_por_id.get(principal, set()):
            r.avisos.append(f"{base}: a verba {principal} não marca reflexos.{chave} = true; o reflexo manual não será criado.")
            continue
        descricao = _descricao(r, ler)
        if not descricao:
            ler.erro("descricao", "obrigatória")
        assunto = item.get("assuntoCnj")
        if assunto in (None, ""):
            assunto = ASSUNTO_CNJ_PADRAO
            r.avisos.append(f"{base}: assuntoCnj não informado; usado {ASSUNTO_CNJ_PADRAO} (o da planilha de verbas da casa).")
        elif not str(assunto).isdigit():
            ler.erro("assuntoCnj", "código numérico do assunto CNJ")
        linha = {"SNMVERBAREFLEXA": REFLEXOS[chave], "SNMDESCRICAOVERBAREFLEXA": descricao, "ICDASSUNTO": int(str(assunto)) if str(assunto).isdigit() else None,
                 "STPTIPOVERBA": "R"}
        for campo in ("bases", "baseTabelada"):
            if campo in item:
                r.ignorar(f"{base}.{campo}", "a base do reflexo manual é sempre a verba principal")
        caracteristica = ler.lista("caracteristica", CARACTERISTICA)
        _, _, ocorrencia = _parametros_comuns(ler, linha)
        if ocorrencia is None and caracteristica:
            ocorrencia = OCORRENCIA_DA_CARACTERISTICA[caracteristica]
            linha["STPOCORRENCIAPAGAMENTO"] = OCORRENCIA[ocorrencia]
            r.avisos.append(f"{base}: ocorrenciaPagamento derivada da característica: {ocorrencia}.")
        comportamento = ler.lista("comportamento", COMPORTAMENTO)
        media = comportamento not in (None, "VALOR_MENSAL")
        if media and ocorrencia in (None, "MENSAL"):
            ler.erro("comportamento", "média não existe com ocorrência MENSAL; mande ocorrenciaPagamento"
                     if ocorrencia is None or item.get("ocorrenciaPagamento") is None else "média não existe com ocorrência MENSAL")
        periodo = item.get("periodoMedia")
        if media and periodo is not None:
            aceitos = PERIODOS_DA_OCORRENCIA.get(ocorrencia, ())
            if periodo not in aceitos:
                ler.erro("periodoMedia", f"fora da lista para a ocorrência {ocorrencia}", aceitos)
                periodo = None
        elif periodo is not None:
            periodo = None  # sem média a tela não mostra o campo; o .PJC grava um padrão — ignorado, nunca recusado
        fracao = ler.lista("tratamentoFracaoMes", FRACAO)
        if fracao in ("DESPREZAR", "DESPREZAR_MENOR_QUE_15_DIAS") and not media:
            ler.erro("tratamentoFracaoMes", "DESPREZAR só com comportamento de média", ["MANTER", "INTEGRALIZAR"])
        linha.update({"STPCOMPORTAMENTOREFLEXO": COMPORTAMENTO.get(comportamento), "STPPERIODOMEDIAREFLEXO": PERIODO_MEDIA.get(periodo),
                      "STPTRATAMENTOFRACAOMESREFLEXO": FRACAO.get(fracao)})
        manuais.append({"principal": principal, "reflexo": chave, "linha": linha})

    # reflexo marcado que exige manual e não tem: recusa na porta (na execução viraria verba criada pela metade)
    tem_manual = {(m["principal"], m["reflexo"]) for m in manuais}
    for ident, marcados in marcados_por_id.items():
        for chave in sorted(marcados - SEMPRE_OFERECIDOS):
            if (ident, chave) not in tem_manual and ident in por_id:
                r.erro(f"{por_id[ident]['base']}.reflexos.{chave}", True,
                       f"o reflexo '{chave}' precisa de um item em reflexosManuais com principal = {ident}")
    return manuais


def criar_planilha_de_verbas(caminho_json, verbas_pedido, execucao_id):
    """Planilha de verbas temporária, ao lado do .json, no layout que a classe Verbas lê."""
    import openpyxl

    destino = Path(caminho_json).with_name(f"{execucao_id}_automacao_verbas.xlsx")
    livro = openpyxl.Workbook()
    livro.active.title = "TBCALCULO"
    livro.active.append(["Pedido montado do .json de lançamentos (bloco verbas)."])

    aba = livro.create_sheet("TBVERBABD")
    aba.append(COLUNAS_VERBA)
    for v in verbas_pedido["verbas"]:
        aba.append([v["linha"].get(c) for c in COLUNAS_VERBA])

    aba = livro.create_sheet("TBREFLEXA")
    aba.append(COLUNAS_REFLEXA)
    for n, m in enumerate(verbas_pedido["reflexos"], start=1):
        m["linha"]["IIDVERBAREFLEXA"] = n
        aba.append([m["linha"].get(c) for c in COLUNAS_REFLEXA])

    aba = livro.create_sheet("DEVIDOBD_EXP")
    aba.append(["TIPO", "PARCELA", "PROP/INT", "FKVERBA"])
    for v in verbas_pedido["verbas"]:
        for tipo, parcela, prop in v["devido"]:
            aba.append([tipo, parcela, prop, v["linha"]["IIDVERBAJRS"]])

    aba = livro.create_sheet("PAGOBD_EXP")
    aba.append(["PARCELA", "PROP", "FKVERBA"])
    for v in verbas_pedido["verbas"]:
        for parcela, prop in v["pago"]:
            aba.append([parcela, prop, v["linha"]["IIDVERBAJRS"]])
    livro.save(destino)
    return str(destino)
