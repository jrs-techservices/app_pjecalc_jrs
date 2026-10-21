# -*- coding: utf-8 -*-
"""Lê o .json de lançamentos do SAP (contrato lancamentos-pjecalc/1) e o traduz para chaves da PJE-BD da v3.41.

As telas continuam lendo a PJE-BD; este módulo só produz as linhas que faltam, no formato que elas já entendem.
Valor fora das listas do PJe-Calc recusa o pedido na porta; campo que a automação não conhece é ignorado e relatado.
"""
import csv
import json
import os
import re
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

from servico import config
from Calculo.leitura_planilha import data_da_planilha, numero_da_planilha, documento_da_planilha, opcao_de_municipio, UFS

CONTRATO = "lancamentos-pjecalc/1"
CATALOGO = json.loads((Path(__file__).with_name("catalogo_pjecalc.json")).read_text(encoding="utf-8"))
ORIGEM_EXCEL = datetime(1899, 12, 30)
PADRAO_CNJ = re.compile(r"^\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4}$")
META = {"contrato", "geradoEm", "procedencia", "processo.numero"}
ORDEM = [
    "id_processo", "numero_processo", "valor_da_causa", "nome_parte1", "doc_parte1", "advogado_parte1", "nome_parte2",
    "doc_parte2", "advogado_parte2", "uf", "municipio", "data_admissao", "data_rescisao", "data_ajuizamento",
    "inicio_calculo", "termino_calculo", "aplicar_prescricao_verbas", "aplicar_prescricao_FGTS", "regime_trabalho",
    "maior_remuneracao", "ultima_remuneracao", "prazo_de_aviso_previo", "quantidade_dias_aviso_previo",
    "projetar_aviso_previo_indenizado", "limitar_avos_ao_periodo_do_calculo", "zerar_valor_negativo_padrao",
    "considerar_feriados_estaduais", "considerar_feriados_municipais", "carga_horaria", "sabado_como_dia_util",
    "comentarios_parametros", "ponto_facultutativo_sextafeirasanta", "ponto_facultutativo_corpuschristi",
    "ponto_facultutativo_carnaval", "destino_fgts", "compor_principal", "multa_fgts", "cont_social_aliquota_empregador",
    "processo_de_agroindustria", "empresa_enquadrada_simples", "empresa_com_desoneracao", "entidade_filantropica",
]
MINIMO_SEM_PLANILHA = {"nome_parte1": "reclamante", "nome_parte2": "reclamado", "uf": "UF", "municipio": "município",
                       "data_admissao": "admissão", "data_ajuizamento": "ajuizamento"}

TIPOS_HONORARIO = {
    "ADVOCATICIOS": "Honorários Advocatícios", "ASSISTENCIAIS": "Honorários Assistenciais",
    "CONTRATUAIS": "Honorários Contratuais", "PERICIAIS_CONTADOR": "Honorários Periciais - Contador",
    "PERICIAIS_DOCUMENTOSCOPIO": "Honorários Periciais - Documentoscópio",
    "PERICIAIS_ENGENHEIRO": "Honorários Periciais - Engenheiro", "PERICIAIS_INTERPRETE": "Honorários Periciais - Intérprete",
    "PERICIAIS_MEDICO": "Honorários Periciais - Médico", "PERICIAIS_OUTROS": "Honorários Periciais - Outros",
    "SUCUMBENCIAIS": "Honorários de Sucumbência", "LEILOEIRO": "Leiloeiro",
}
CREDOR_DEVEDOR_MULTA = {
    "RECLAMANTE_RECLAMADO": "Reclamante e Reclamado", "RECLAMADO_RECLAMANTE": "Reclamado e Reclamante",
    "TERCEIRO_RECLAMANTE": "Terceiro e Reclamante", "TERCEIRO_RECLAMADO": "Terceiro e Reclamado",
    "RECLAMANTE_E_RECLAMADO": "Reclamante e Reclamado", "RECLAMADO_E_RECLAMANTE": "Reclamado e Reclamante",
    "TERCEIRO_E_RECLAMANTE": "Terceiro e Reclamante", "TERCEIRO_E_RECLAMADO": "Terceiro e Reclamado",
}
CUSTAS_RECLAMADO = {"NAO_SE_APLICA": "Não se Aplica", "CALCULADA_2_POR_CENTO": "Calculada 2%", "INFORMADA": "Informada"}
DESTINO_FGTS = {"RECOLHER": "Recolher", "PAGAR": "Pagar"}
DEVEDOR = {"RECLAMADO": "RECLAMADA", "RECLAMANTE": "RECLAMANTE"}
COBRANCA = {"COBRAR": "COBRAR", "DESCONTAR": "DESCONTAR", "DESCONTAR_CREDITO": "DESCONTAR"}


class LancamentosIlegiveis(ValueError):
    """O arquivo não pôde ser lido como o contrato esperado; o pedido é recusado inteiro."""


def _sem_acento(texto):
    return unicodedata.normalize("NFKD", str(texto)).encode("ASCII", "ignore").decode("ASCII").lower()


def _serial(valor):
    data = data_da_planilha(valor)
    return (data - ORIGEM_EXCEL).days if data else None


def _texto(valor):
    return "True" if valor else "False"


def _sim_nao(valor):
    return "Sim" if valor else "Não"


# caminho do .json → tela do PJe-Calc (para relatar os ignorados tela a tela; a ordem importa: o 1º prefixo que casar vale)
TELA_DO_CAMINHO = [
    (("dadosDoCalculo", "telas.dados-do-calculo", "advogados", "parametrosDoCalculo", "processo"), "Dados do Cálculo"),
    (("arquivos.faltas",), "Faltas"), (("arquivos.ferias",), "Férias"), (("arquivos.historico-salarial",), "Histórico Salarial"),
    (("arquivos.cartao-de-ponto", "telas.cartao-de-ponto", "cartaoDePonto"), "Cartão de Ponto"),
    (("verbas", "reflexosManuais"), "Verbas"),
    (("telas.salario-familia",), "Salário-família"), (("telas.seguro-desemprego", "seguroDesemprego"), "Seguro-desemprego"),
    (("fgts", "telas.fgts", "fgtsSaldoSaque"), "FGTS"),
    (("telas.contribuicao-social", "contribuicaoSocial", "arquivos.base-inss"), "Contribuição Social"),
    (("telas.previdencia-privada", "previdenciaPrivada"), "Previdência Privada"), (("telas.pensao-alimenticia",), "Pensão Alimentícia"),
    (("telas.imposto-de-renda",), "Imposto de Renda"),
    (("multas", "telas.multas"), "Multas e Indenizações"), (("honorarios", "telas.honorarios"), "Honorários"),
    (("custas", "telas.custas-judiciais"), "Custas Judiciais"), (("telas.correcao-juros-multa",), "Correção, Juros e Multa"),
    (("dataLiquidacao",), "Liquidar"),
]


def tela_do_caminho(caminho):
    for prefixos, tela in TELA_DO_CAMINHO:
        if any(caminho == p or caminho.startswith((p + ".", p + "[")) for p in prefixos):
            return tela
    return "Outros"


def ignorados_por_tela(ignorados):
    """{tela: {campo: motivo}} na ordem do fluxo do robô — para a conferência saber que "não conferível" é contrato pendente."""
    ordem = [t for _, t in TELA_DO_CAMINHO] + ["Outros"]
    agrupado = {}
    for campo, motivo in sorted(ignorados.items()):
        agrupado.setdefault(tela_do_caminho(campo), {})[campo] = motivo
    return {t: agrupado[t] for t in ordem if t in agrupado}


class Interpretacao:

    def __init__(self, dados):
        self.dados = dados
        self.valores = {}
        self.familias = []
        self.usados = set()
        self.ignorados = {}
        self.avisos = []
        self.erros = []
        self.faltas = []
        self.ferias = []
        self.historicos = []
        self.base_inss = []
        self.arquivos = {}

    # --- registro ---
    def usar(self, caminho, chave, valor):
        if valor is None:
            return
        self.valores[chave] = valor
        self.usados.add(caminho)

    def ignorar(self, caminho, motivo):
        self.ignorados.setdefault(caminho, motivo)

    def erro(self, caminho, valor, mensagem, aceitos=None):
        item = {"campo": caminho, "valor": valor, "mensagem": mensagem}
        if aceitos:
            item["aceitos"] = list(aceitos)
        self.erros.append(item)

    def enumeracao(self, caminho, valor, tabela):
        """Código do .json → texto da tela; aceita também o próprio texto da tela."""
        if valor in tabela:
            return tabela[valor]
        if valor in tabela.values():
            return valor
        self.erro(caminho, valor, "valor não reconhecido", sorted(tabela))
        return None

    def data(self, caminho, valor):
        serial = _serial(valor)
        if serial is None:
            self.erro(caminho, valor, "data inválida: use AAAA-MM-DD")
        return serial

    # --- resultado ---
    def pares_avulsos(self):
        """(chave, valor) na ordem da v3.41: há chaves de disparo que precisam vir depois das outras."""
        ordenados = [(c, self.valores[c]) for c in ORDEM if c in self.valores]
        return ordenados + [(c, v) for c, v in self.valores.items() if c not in ORDEM]

    def pares(self):
        return self.pares_avulsos() + [par for _, familia in self.familias for par in familia]

    def faltando_para_criar(self):
        return [rotulo for chave, rotulo in MINIMO_SEM_PLANILHA.items() if chave not in self.valores]

    def resumo(self):
        return {
            "contrato": self.dados.get("contrato"),
            "usados": sorted(self.usados),
            "ignorados": dict(sorted(self.ignorados.items())),
            "ignorados_por_tela": ignorados_por_tela(self.ignorados),
            "avisos": self.avisos,
        }


def _folhas(dados, prefixo=""):
    """Caminhos de todos os campos do .json, com listas como 'bloco[]'."""
    if isinstance(dados, dict):
        for chave, valor in dados.items():
            yield from _folhas(valor, f"{prefixo}.{chave}" if prefixo else chave)
    elif isinstance(dados, list) and dados and all(isinstance(i, dict) for i in dados):
        for item in dados:
            yield from _folhas(item, f"{prefixo}[]")
    else:
        yield prefixo


# ---------------------------------------------------------------- blocos do documento anterior (dadosDoCalculo, …)

def _dados_do_calculo(r):
    bloco = r.dados.get("dadosDoCalculo") or {}
    processo = r.dados.get("processo") or {}
    if processo.get("id") not in (None, ""):
        r.usar("processo.id", "id_processo", int(processo["id"]) if str(processo["id"]).isdigit() else processo["id"])
    numero = bloco.get("numeroProcesso") or processo.get("numero")
    if numero:
        if PADRAO_CNJ.match(str(numero)):
            r.usar("dadosDoCalculo.numeroProcesso", "numero_processo", str(numero))
        else:
            r.erro("dadosDoCalculo.numeroProcesso", numero, "número CNJ fora do formato 0000000-00.0000.0.00.0000")
    vistos = {}
    for i, parte in enumerate(bloco.get("partes") or []):
        polo = str(parte.get("polo", "")).upper()
        sufixo = {"ATIVO": "1", "PASSIVO": "2"}.get(polo)
        if not sufixo:
            r.erro(f"dadosDoCalculo.partes[{i}].polo", parte.get("polo"), "polo não reconhecido", ["ATIVO", "PASSIVO"])
            continue
        if polo in vistos:
            # o PJe-Calc tem um reclamante e um reclamado: o 2º do mesmo polo apagaria o 1º sem aviso
            r.erro(f"dadosDoCalculo.partes[{i}].polo", parte.get("polo"),
                   f"o PJe-Calc aceita uma parte por polo; o polo {polo} já veio em partes[{vistos[polo]}]")
            continue
        vistos[polo] = i
        if parte.get("nome"):
            r.usar("dadosDoCalculo.partes[].nome", f"nome_parte{sufixo}", str(parte["nome"])[:150])
        if parte.get("documento"):
            r.usar("dadosDoCalculo.partes[].documento", f"doc_parte{sufixo}", documento_da_planilha(parte["documento"]))
        else:
            r.usados.add("dadosDoCalculo.partes[].documento")
        r.usados.add("dadosDoCalculo.partes[].polo")
    for campo, chave in (("dataAdmissao", "data_admissao"), ("dataDemissao", "data_rescisao"),
                         ("dataAjuizamento", "data_ajuizamento"), ("dataInicioCalculo", "inicio_calculo"),
                         ("dataTerminoCalculo", "termino_calculo")):
        if bloco.get(campo):
            r.usar(f"dadosDoCalculo.{campo}", chave, r.data(f"dadosDoCalculo.{campo}", bloco[campo]))
    if bloco.get("atividadeEconomica"):
        r.usar("dadosDoCalculo.atividadeEconomica", "cont_social_aliquota_empregador", str(bloco["atividadeEconomica"]).strip())
    if bloco.get("vara"):
        r.usados.add("dadosDoCalculo.vara")
        _municipio_da_vara(r, str(bloco["vara"]))
    for campo, motivo in (("identificacao", "só identifica o processo no SAP"), ("classe", "sem campo no PJe-Calc"),
                          ("tipoRescisao", "sem campo na tela Dados do Cálculo")):
        if campo in bloco:
            r.ignorar(f"dadosDoCalculo.{campo}", motivo)


def _municipio_da_vara(r, vara):
    """'19ª VARA DO TRABALHO DO RECIFE-PE' → RECIFE / PE, conferido na lista do PJe-Calc."""
    telas_dados = (r.dados.get("telas") or {}).get("dados-do-calculo") or {}
    if "municipio" in r.valores or telas_dados.get("municipio"):
        return
    achado = re.search(r"^(.*?)\s*[-/]\s*([A-Za-z]{2})\s*$", vara.strip())
    if not achado or achado.group(2).upper() not in UFS:
        r.avisos.append(f"Município não identificado na vara '{vara}'.")
        return
    uf = achado.group(2).upper()
    palavras = achado.group(1).split()
    for inicio in range(len(palavras)):
        candidato = " ".join(palavras[inicio:])
        opcao = opcao_de_municipio(candidato, CATALOGO["municipios"].get(uf, []))
        if opcao and _sem_acento(opcao) == _sem_acento(" ".join(candidato.upper().split())):
            r.usar("dadosDoCalculo.vara", "uf", uf)
            r.usar("dadosDoCalculo.vara", "municipio", opcao)
            r.avisos.append(f"Município e UF tirados da vara: {opcao}/{uf}. Mande 'municipio' e 'uf' para não depender disso.")
            return
    r.avisos.append(f"Município não identificado na vara '{vara}'.")


def _fgts(r):
    bloco = r.dados.get("fgts") or {}
    if "destino" in bloco:
        r.usar("fgts.destino", "destino_fgts", r.enumeracao("fgts.destino", bloco["destino"], DESTINO_FGTS))
    if "multaDoFgts" in bloco:
        r.usar("fgts.multaDoFgts", "multa_fgts", _sim_nao(bool(bloco["multaDoFgts"])))


def _codigo(texto):
    """'Principal (-) Contribuição Social' → 'PRINCIPAL_MENOS_CONTRIBUICAO_SOCIAL' (a forma dos códigos do XML)."""
    texto = _sem_acento(texto).upper().replace("(-)", " MENOS ")
    return re.sub(r"[^A-Z0-9]+", "_", texto).strip("_")


def _opcao_da_tela(r, caminho, valor, opcoes):
    """Aceita o texto da lista do PJe-Calc ou o código equivalente (com ou sem o '_MENOS_')."""
    if valor in opcoes:
        return valor
    alvo = _codigo(valor or "")
    for opcao in opcoes:
        codigo = _codigo(opcao)
        if alvo and alvo in (codigo, codigo.replace("_MENOS", "")):
            return opcao
    r.erro(caminho, valor, "valor não reconhecido", opcoes)
    return None


def _fracao(r, caminho, valor):
    numero = numero_da_planilha(valor)
    if numero is None or not 0 < numero <= 1:
        r.erro(caminho, valor, "alíquota em FRAÇÃO, entre 0 e 1 (0.05 = 5%)")
        return None
    return float(numero)


TIPOS_CUSTAS_RECLAMANTE = ("NAO_SE_APLICA", "CALCULADA_2_POR_CENTO", "INFORMADA")
TIPOS_CUSTAS_LIQUIDACAO = ("NAO_SE_APLICA", "CALCULADA_MEIO_POR_CENTO", "INFORMADA")
BASES_CUSTAS = ("BRUTO_DEVIDO_AO_RECLAMANTE", "BRUTO_DEVIDO_AO_RECLAMANTE_MAIS_DEBITOS_RECLAMADO")


def _custas_completas(r, bloco):
    """Campos da tela Custas Judiciais além das do reclamado: base das calculadas, custas do reclamante (tipo, cobrança,
    vencimento/valor se INFORMADA), de liquidação (tipo, vencimento/valor se INFORMADA) e as recolhidas pelo reclamante.
    Códigos = os do .PJC. Ausente = padrão da tela."""
    def codigo(campo, opcoes, chave):
        valor = bloco.get(campo)
        if valor in (None, ""):
            return None
        if valor not in opcoes:
            r.erro(f"custas.{campo}", valor, "valor não reconhecido", opcoes)
            return None
        r.usar(f"custas.{campo}", chave, valor)
        return valor

    def informada(tipo, campo_data, campo_valor, chave_data, chave_valor, rotulo):
        for campo in (campo_data, campo_valor):
            if campo in bloco and tipo != "INFORMADA":
                r.erro(f"custas.{campo}", bloco[campo], f"só com {rotulo} = INFORMADA")
        if tipo != "INFORMADA":
            return
        serial = r.data(f"custas.{campo_data}", bloco.get(campo_data))
        valor = numero_da_planilha(bloco.get(campo_valor))
        if valor is None or valor <= 0:
            r.erro(f"custas.{campo_valor}", bloco.get(campo_valor), f"valor obrigatório (> 0) com {rotulo} = INFORMADA")
        if serial is not None and valor is not None and valor > 0:
            r.usar(f"custas.{campo_data}", chave_data, f"{ORIGEM_EXCEL + timedelta(days=serial):%d/%m/%Y}")
            r.usar(f"custas.{campo_valor}", chave_valor, float(valor))

    codigo("baseParaCustasCalculadas", BASES_CUSTAS, "custas_base")
    reclamante = codigo("tipoDeCustasDeConhecimentoDoReclamante", TIPOS_CUSTAS_RECLAMANTE, "custas_rcte_tipo")
    if "tipoCobrancaReclamante" in bloco:
        if reclamante in (None, "NAO_SE_APLICA"):
            r.erro("custas.tipoCobrancaReclamante", bloco["tipoCobrancaReclamante"],
                   "só com tipoDeCustasDeConhecimentoDoReclamante diferente de NAO_SE_APLICA (a tela só mostra assim)")
        else:
            codigo("tipoCobrancaReclamante", ("DESCONTAR_CREDITO", "COBRAR"), "custas_rcte_cobranca")
    informada(reclamante, "dataVencimentoConhecimentoDoReclamante", "valorDeConhecimentoDoReclamante", "custas_rcte_vcto",
              "custas_rcte_valor", "tipoDeCustasDeConhecimentoDoReclamante")
    liquidacao = codigo("tipoDeCustasDeLiquidacao", TIPOS_CUSTAS_LIQUIDACAO, "custas_liq_tipo")
    informada(liquidacao, "dataVencimentoCustasDeLiquidacao", "valorCustasDeLiquidacao", "custas_liq_vcto", "custas_liq_valor",
              "tipoDeCustasDeLiquidacao")
    pagas = bloco.get("custasPagasDoReclamante")
    if pagas not in (None, []):
        if not isinstance(pagas, list) or not all(isinstance(i, dict) for i in pagas):
            r.erro("custas.custasPagasDoReclamante", pagas, "deve ser uma lista de { dataVencimento, valor }")
            return
        linhas, erros_antes = [], len(r.erros)
        for n, paga in enumerate(pagas):
            onde = f"custas.custasPagasDoReclamante[{n}]"
            serial = r.data(f"{onde}.dataVencimento", paga.get("dataVencimento"))
            valor = numero_da_planilha(paga.get("valor"))
            ajuizamento = r.valores.get("data_ajuizamento")
            if valor is None or valor <= 0:
                r.erro(f"{onde}.valor", paga.get("valor"), "valor em reais maior que zero")
            elif serial is not None and isinstance(ajuizamento, int) and serial < ajuizamento:
                # a tela recusa: "O valor do campo Vencimento deve ser igual ou maior que Data de Ajuizamento no Cálculo"
                r.erro(f"{onde}.dataVencimento", paga.get("dataVencimento"), "antes do ajuizamento (a tela recusa)")
            elif serial is not None:
                linhas.append([f"{ORIGEM_EXCEL + timedelta(days=serial):%d/%m/%Y}", float(valor)])
        if len(r.erros) == erros_antes:
            r.valores["custas_pagas_rcte"] = json.dumps(linhas)
            r.usados.update({"custas.custasPagasDoReclamante[].dataVencimento", "custas.custasPagasDoReclamante[].valor"})


def _custas(r):
    bloco = r.dados.get("custas") or {}
    if not bloco:
        return
    _custas_completas(r, bloco)
    tipo = None
    if "tipoDeCustasDeConhecimentoDoReclamado" in bloco:
        tipo = r.enumeracao("custas.tipoDeCustasDeConhecimentoDoReclamado", bloco["tipoDeCustasDeConhecimentoDoReclamado"], CUSTAS_RECLAMADO)
    vencimento_informado, valor_informado = "-", "-"
    if tipo == "Informada":
        vencimento_informado = r.data("custas.dataVencimentoConhecimentoDoReclamado", bloco.get("dataVencimentoConhecimentoDoReclamado"))
        valor_informado = numero_da_planilha(bloco.get("valorConhecimentoDoReclamado"))
        if valor_informado is None:
            r.erro("custas.valorConhecimentoDoReclamado", bloco.get("valorConhecimentoDoReclamado"), "valor obrigatório em custas INFORMADA")
        if None in (vencimento_informado, valor_informado):
            tipo = None
        else:
            r.usados.update({"custas.dataVencimentoConhecimentoDoReclamado", "custas.valorConhecimentoDoReclamado"})
    pagas = bloco.get("custasPagasDoReclamado") or []
    if tipo is None and not pagas:
        return
    familia = []
    if tipo:
        r.usados.add("custas.tipoDeCustasDeConhecimentoDoReclamado")
        familia += [("custas_reclamado_conhecimento", tipo), ("custas_recda_conhecimento_info_vcto", vencimento_informado),
                    ("custas_recda_conhecimento_info_valor", valor_informado)]
    familia.append(("qtd_custas_recolhidas", len(pagas)))
    for n, paga in enumerate(pagas, start=1):
        valor = numero_da_planilha(paga.get("valor"))
        vencimento = _serial(paga.get("dataVencimento") or paga.get("vencimento"))
        r.usados.add("custas.custasPagasDoReclamado[].valor")
        if vencimento is None:
            r.avisos.append(f"Custa recolhida nº {n} (R$ {valor}) sem data de vencimento: não será lançada. Mande 'dataVencimento'.")
        else:
            r.usados.add("custas.custasPagasDoReclamado[].dataVencimento")
        familia += [(f"custas_recolhidas_rcdo{n}_vencimento", vencimento if vencimento is not None else "-"),
                    (f"custas_recolhidas_rcdo{n}_valor", valor if valor is not None and vencimento is not None else "-")]
    r.familias.append(("custas", familia))


# campos que a tela do honorário mostra em cada tipo de valor (o resto do item não tem onde entrar)
CAMPOS_HONORARIO = {
    "INFORMADO": {"descricao", "tipoHonorario", "tipoDeDevedor", "tipoCobrancaReclamante", "nomeCredor", "tipoValor", "dataVencimento", "valor",
                  "tipoDocumentoFiscalCredor", "numeroDocumentoFiscalCredor", "apurarIRRF", "tipoDeImpostoDeRenda", "apurarIRPFSobreJuros",
                  "aplicarJuros", "dataJurosAPartirDe", "tipoDeIndiceDeCorrecao", "outroIndiceDeCorrecao"},
    "CALCULADO": {"descricao", "tipoHonorario", "tipoDeDevedor", "tipoCobrancaReclamante", "nomeCredor", "tipoValor", "aliquota",
                  "baseParaApuracao", "tipoDeIndiceDeCorrecao", "tipoDocumentoFiscalCredor", "numeroDocumentoFiscalCredor", "apurarIRRF",
                  "tipoDeImpostoDeRenda"},
}


def _extras_do_honorario(r, base, item, tipo_valor):
    """Documento e IR do credor (os dois tipos) e, no INFORMADO, juros e correção por outro índice. Devolve o dicionário
    que a tela aplica (só o que veio) ou None se algo foi recusado. Campo fora da condição em que a tela o mostra → 400."""
    extras, erros_antes = {}, len(r.erros)

    def booleano(campo):
        if campo not in item:
            return None
        if not isinstance(item[campo], bool):
            r.erro(f"{base}.{campo}", item[campo], "use true ou false")
            return None
        return item[campo]

    tipo_doc = item.get("tipoDocumentoFiscalCredor")
    numero = re.sub(r"\D", "", str(item.get("numeroDocumentoFiscalCredor") or ""))
    if tipo_doc not in (None, "") or numero:
        tamanhos = {"CPF": 11, "CNPJ": 14, "CEI": 12}
        if tipo_doc not in tamanhos:
            r.erro(f"{base}.tipoDocumentoFiscalCredor", tipo_doc, "valor não reconhecido (vem junto com o número)", list(tamanhos))
        elif len(numero) != tamanhos[tipo_doc] or (tipo_doc == "CPF" and not cpf_valido(numero)) or (tipo_doc == "CNPJ" and not cnpj_valido(numero)):
            r.erro(f"{base}.numeroDocumentoFiscalCredor", item.get("numeroDocumentoFiscalCredor"), f"{tipo_doc} inválido")
        else:
            extras["doc_tipo"], extras["doc_numero"] = tipo_doc, numero
    irrf = booleano("apurarIRRF")
    if irrf is not None:
        if irrf and "doc_numero" not in extras:
            r.erro(f"{base}.apurarIRRF", True, "apurar IR exige o documento do credor (tipoDocumentoFiscalCredor + numeroDocumentoFiscalCredor)")
        extras["irrf"] = irrf
    if item.get("tipoDeImpostoDeRenda") not in (None, ""):
        if not irrf:
            r.erro(f"{base}.tipoDeImpostoDeRenda", item["tipoDeImpostoDeRenda"], "só com apurarIRRF = true")
        elif item["tipoDeImpostoDeRenda"] not in ("PESSOA_FISICA", "PESSOA_JURIDICA"):
            r.erro(f"{base}.tipoDeImpostoDeRenda", item["tipoDeImpostoDeRenda"], "valor não reconhecido", ["PESSOA_FISICA", "PESSOA_JURIDICA"])
        else:
            extras["ir_tipo"] = item["tipoDeImpostoDeRenda"]
    if tipo_valor == "INFORMADO":
        juros = booleano("aplicarJuros")
        if juros is not None:
            extras["juros"] = juros
            if juros:
                desde = r.data(f"{base}.dataJurosAPartirDe", item.get("dataJurosAPartirDe"))
                if desde is not None:
                    extras["juros_desde"] = f"{ORIGEM_EXCEL + timedelta(days=desde):%d/%m/%Y}"
        if "dataJurosAPartirDe" in item and not juros:
            r.erro(f"{base}.dataJurosAPartirDe", item["dataJurosAPartirDe"], "só com aplicarJuros = true")
        ir_juros = booleano("apurarIRPFSobreJuros")
        if ir_juros is not None:
            if ir_juros and not (juros and irrf):
                r.erro(f"{base}.apurarIRPFSobreJuros", True, "só com aplicarJuros = true e apurarIRRF = true (a tela só habilita assim)")
            extras["ir_juros"] = ir_juros
        indice = item.get("tipoDeIndiceDeCorrecao")
        if indice not in (None, "", "UTILIZAR_INDICE_TRABALHISTA"):
            outro = item.get("outroIndiceDeCorrecao")
            aceitos = CATALOGO["telas"]["correcao-juros-multa"]["listas"]["indiceTrabalhista"]
            if indice != "UTILIZAR_OUTRO_INDICE":
                r.erro(f"{base}.tipoDeIndiceDeCorrecao", indice, "valor não reconhecido", ["UTILIZAR_INDICE_TRABALHISTA", "UTILIZAR_OUTRO_INDICE"])
            elif outro not in aceitos:
                r.erro(f"{base}.outroIndiceDeCorrecao", outro, "obrigatório com UTILIZAR_OUTRO_INDICE; texto da lista indiceTrabalhista", aceitos)
            else:
                extras["indice_outro"] = outro
    return extras if len(r.erros) == erros_antes else None

# Honorário CALCULADO vai pela família 'contratuais' da tela: ela pula tipo com '-' (os periciais) e aplica .title()
# ao tipo (o que estraga 'Honorários de Sucumbência'); por isso só estes quatro.
TIPOS_HONORARIO_CALCULADO = {"Honorários Advocatícios", "Honorários Assistenciais", "Honorários Contratuais", "Leiloeiro"}


def _usar_campos_do_honorario(r, n, item, tipo_valor):
    for campo in item:
        if campo in CAMPOS_HONORARIO[tipo_valor]:
            r.usados.add(f"honorarios[].{campo}")
        elif campo == "dataVencimento":
            r.ignorar(f"honorarios[{n}].dataVencimento", "no honorário CALCULADO a tela não tem vencimento: o PJe-Calc usa a data da liquidação")
        else:
            r.ignorar(f"honorarios[{n}].{campo}", f"a tela do honorário {tipo_valor} não tem este campo")


def _latin1(r, caminho, texto):
    """O PJe-Calc grava texto em Latin-1 (travessão e aspas curvas viram '¿'): troca pelo equivalente simples, com aviso."""
    from servico.verbas_pedido import texto_latin1
    limpo = texto_latin1(texto)
    if limpo != texto:
        r.avisos.append(f"{caminho}: caracteres que o PJe-Calc não grava foram trocados: {texto!r} -> {limpo!r}")
    return limpo


def _honorarios(r):
    itens = r.dados.get("honorarios")
    if itens is None:
        return
    if itens == []:
        r.usados.add("honorarios")  # lista vazia é a decisão "não há honorarios": a tela fica sem lançamentos
    informados, calculados = [], []
    for n, item in enumerate(itens):
        base = f"honorarios[{n}]"
        tipo_valor = str(item.get("tipoValor", "INFORMADO")).upper()
        if tipo_valor not in ("INFORMADO", "CALCULADO"):
            r.erro(f"{base}.tipoValor", item.get("tipoValor"), "valor não reconhecido", ["INFORMADO", "CALCULADO"])
            continue
        descricao = _latin1(r, f"{base}.descricao", str(item.get("descricao") or "").strip())
        if len(descricao) > 50 and tipo_valor == "INFORMADO":
            r.avisos.append(f"{base}.descricao: tem {len(descricao)} caracteres; no honorário INFORMADO o PJe-Calc guarda os 50 "
                            f"primeiros ({descricao[:50].upper()!r})")
        if not descricao or "-" in descricao:
            r.erro(f"{base}.descricao", item.get("descricao"), "descrição obrigatória e sem '-' nem travessão (a automação usa '-' "
                                                                "para 'bloco vazio')")
            continue
        tipo = r.enumeracao(f"{base}.tipoHonorario", item.get("tipoHonorario"), TIPOS_HONORARIO)
        devedor = r.enumeracao(f"{base}.tipoDeDevedor", item.get("tipoDeDevedor"), DEVEDOR)
        cobranca = r.enumeracao(f"{base}.tipoCobrancaReclamante", item.get("tipoCobrancaReclamante", "COBRAR"), COBRANCA)
        credor = _latin1(r, f"{base}.nomeCredor", str(item.get("nomeCredor") or ""))
        extras = _extras_do_honorario(r, base, item, tipo_valor)
        if extras is None:
            continue
        if tipo_valor == "INFORMADO":
            vencimento = r.data(f"{base}.dataVencimento", item.get("dataVencimento"))
            valor = numero_da_planilha(item.get("valor"))
            if valor is None:
                r.erro(f"{base}.valor", item.get("valor"), "valor obrigatório em honorário INFORMADO")
            if None not in (tipo, devedor, cobranca, vencimento, valor):
                informados.append((tipo, descricao, devedor, cobranca, vencimento, valor, credor, extras))
                _usar_campos_do_honorario(r, n, item, "INFORMADO")
            continue
        if tipo and tipo not in TIPOS_HONORARIO_CALCULADO:
            r.erro(f"{base}.tipoHonorario", item.get("tipoHonorario"), "honorário CALCULADO deste tipo ainda não suportado",
                   sorted(k for k, v in TIPOS_HONORARIO.items() if v in TIPOS_HONORARIO_CALCULADO))
            continue
        if item.get("outroIndiceDeCorrecao") or "OUTRO" in _codigo(item.get("tipoDeIndiceDeCorrecao") or ""):
            r.erro(f"{base}.outroIndiceDeCorrecao", item.get("outroIndiceDeCorrecao") or item.get("tipoDeIndiceDeCorrecao"),
                   "índice de correção próprio em honorário CALCULADO ainda não suportado (sai o índice trabalhista)")
            continue
        aliquota = _fracao(r, f"{base}.aliquota", item.get("aliquota"))
        base_apuracao = _opcao_da_tela(r, f"{base}.baseParaApuracao", item.get("baseParaApuracao"),
                                       CATALOGO["telas"]["honorarios"]["listas"]["baseParaApuracao"])
        if None not in (tipo, devedor, cobranca, aliquota, base_apuracao):
            calculados.append((tipo, descricao, devedor, "DEDUZIR" if cobranca == "DESCONTAR" else "COBRAR", aliquota, base_apuracao, credor, extras))
            _usar_campos_do_honorario(r, n, item, "CALCULADO")

    familia = [("qtd_honorarios_periciais", len(informados))]
    for n, (tipo, descricao, devedor, cobranca, vencimento, valor, credor, extras) in enumerate(informados, start=1):
        if extras:
            familia.append((f"honorarios_periciais_extras_{n}", json.dumps(extras, ensure_ascii=False)))
        familia += [(f"honorarios_periciais_tipo_pericia_pje_calc_{n}", tipo), (f"honorarios_periciais_descricao_{n}", descricao),
                    (f"honorarios_periciais_devedor_{n}", devedor), (f"honorarios_periciais_pgto_rcte_{n}", cobranca),
                    (f"honorarios_periciais_data_{n}", vencimento), (f"honorarios_periciais_valor_{n}", valor),
                    (f"honorarios_periciais_nome_credor_{n}", credor), (f"honorarios_periciais_cm_{n}", "<branco>")]
    familia.append(("qtd_honorarios_contratuais", len(calculados)))
    for n, (tipo, descricao, devedor, participacao, aliquota, base_apuracao, credor, extras) in enumerate(calculados, start=1):
        if extras:
            familia.append((f"honorarios_contratuais_extras_{n}", json.dumps(extras, ensure_ascii=False)))
        familia += [(f"honorarios_contratuais_tipob_{n}", tipo), (f"honorarios_contratuais_descricao_{n}", descricao),
                    (f"honorarios_contratuais_devedor_{n}", devedor), (f"honorarios_contratuais_partic_{n}", participacao),
                    (f"honorarios_contratuais_tipo_{n}", "CALCULADO"), (f"honorarios_contratuais_alíquota_{n}", aliquota),
                    (f"honorarios_contratuais_base_{n}", base_apuracao), (f"honorarios_contratuais_nome_{n}", credor)]
    r.familias.append(("honorarios", familia))


def _usar_campos_da_multa(r, n, item, tipo, credor, juros):
    """Só conta como usado o que a tela da multa aplica neste caso; o resto vai para ignorados, por item, com o motivo
    (antes, qualquer campo do item virava "usado", inclusive os que a automação não digita)."""
    aplicados = {"descricao", "tipoCredorDevedor", "tipoValorDaMulta"}
    if credor and credor.startswith("Terceiro"):
        aplicados.add("nomeTerceiro")
    if credor == "Terceiro e Reclamante":
        aplicados.add("tipoCobrancaReclamante")
    if tipo == "INFORMADO":
        aplicados |= {"dataVencimentoMulta", "valorMulta", "aplicarJurosSobreMulta", "opcaoIndiceDeCorrecaoDaMulta",
                      "outroIndiceDeCorrecaoDaMulta"} | ({"dataJurosAPartirDe"} if juros else set())
    else:
        aplicados |= {"aliquotaMulta", "tipoBaseMulta", "aplicarJurosSobreMulta"}
    for campo in item:
        if campo in aplicados:
            r.usados.add(f"multas[].{campo}")
        elif campo == "tipoCobrancaReclamante":
            r.ignorar(f"multas[{n}].{campo}", "a tela só pergunta a cobrança com credor TERCEIRO_RECLAMANTE")
        elif campo in ("opcaoIndiceDeCorrecaoDaMulta", "outroIndiceDeCorrecaoDaMulta"):
            r.ignorar(f"multas[{n}].{campo}", "correção própria só existe na multa INFORMADA (a CALCULADA usa o índice trabalhista)")
        else:
            r.ignorar(f"multas[{n}].{campo}", f"a tela da multa {tipo} não usa este campo neste caso")


def _multas(r):
    itens = r.dados.get("multas")
    if itens is None:
        return
    if itens == []:
        r.usados.add("multas")  # lista vazia é a decisão "não há multas": a tela fica sem lançamentos
    if len(itens) > 9:
        r.erro("multas", len(itens), "no máximo 9 multas por pedido (limitação da leitura por nome de chave)")
        return
    familia = [("multas_indenizacoes_qtd", len(itens))]
    for n, item in enumerate(itens, start=1):
        base = f"multas[{n - 1}]"
        tipo = str(item.get("tipoValorDaMulta", "INFORMADO")).upper()
        if tipo not in ("INFORMADO", "CALCULADO"):
            r.erro(f"{base}.tipoValorDaMulta", item.get("tipoValorDaMulta"), "valor não reconhecido", ["INFORMADO", "CALCULADO"])
            continue
        credor = r.enumeracao(f"{base}.tipoCredorDevedor", item.get("tipoCredorDevedor"), CREDOR_DEVEDOR_MULTA)
        terceiro = "-"
        if credor and credor.startswith("Terceiro"):
            terceiro = _latin1(r, f"{base}.nomeTerceiro", str(item.get("nomeTerceiro") or "").strip()) or None
            if terceiro is None:
                r.erro(f"{base}.nomeTerceiro", item.get("nomeTerceiro"), "nome do terceiro obrigatório quando o credor/devedor é um terceiro")
        cobranca = r.enumeracao(f"{base}.tipoCobrancaReclamante", item.get("tipoCobrancaReclamante", "COBRAR"), COBRANCA)
        descricao = _latin1(r, f"{base}.descricao", str(item.get("descricao") or "").strip())
        if not descricao:
            r.erro(f"{base}.descricao", item.get("descricao"), "descrição obrigatória")
        vencimento, valor, base_multa, aliquota, juros_desde = "-", "-", "-", "-", "-"
        indice_outro = None
        juros = bool(item.get("aplicarJurosSobreMulta"))
        if tipo == "INFORMADO":
            vencimento = r.data(f"{base}.dataVencimentoMulta", item.get("dataVencimentoMulta"))
            valor = numero_da_planilha(item.get("valorMulta"))
            if valor is None:
                r.erro(f"{base}.valorMulta", item.get("valorMulta"), "valor obrigatório em multa INFORMADA")
            if juros:
                juros_desde = _serial(item.get("dataJurosAPartirDe"))
                if juros_desde is None:
                    r.erro(f"{base}.aplicarJurosSobreMulta", True, "com juros, mande 'dataJurosAPartirDe' (AAAA-MM-DD)")
            opcao = item.get("opcaoIndiceDeCorrecaoDaMulta")
            if opcao not in (None, "", "UTILIZAR_INDICE_TRABALHISTA"):
                aceitos = CATALOGO["telas"]["correcao-juros-multa"]["listas"]["indiceTrabalhista"]
                if opcao != "UTILIZAR_OUTRO_INDICE":
                    r.erro(f"{base}.opcaoIndiceDeCorrecaoDaMulta", opcao, "valor não reconhecido", ["UTILIZAR_INDICE_TRABALHISTA", "UTILIZAR_OUTRO_INDICE"])
                elif item.get("outroIndiceDeCorrecaoDaMulta") not in aceitos:
                    r.erro(f"{base}.outroIndiceDeCorrecaoDaMulta", item.get("outroIndiceDeCorrecaoDaMulta"),
                           "obrigatório com UTILIZAR_OUTRO_INDICE; texto da lista indiceTrabalhista", aceitos)
                else:
                    indice_outro = item["outroIndiceDeCorrecaoDaMulta"]
        else:
            if juros:
                r.erro(f"{base}.aplicarJurosSobreMulta", True, "juros em multa CALCULADA ainda não suportados")
                juros_desde = None
            aliquota = _fracao(r, f"{base}.aliquotaMulta", item.get("aliquotaMulta"))
            base_multa = _opcao_da_tela(r, f"{base}.tipoBaseMulta", item.get("tipoBaseMulta"),
                                        CATALOGO["telas"]["multas-e-indenizacoes"]["listas"]["tipoBaseMulta"])
        if None in (credor, terceiro, cobranca, vencimento, valor, juros_desde, aliquota, base_multa) or not descricao:
            continue
        familia += [(f"multas_indenizacoes_descricao_{n}", descricao), (f"multas_indenizacoes_credor_{n}", credor),
                    (f"multas_indenizacoes_form_pagto_{n}", cobranca), (f"multas_indenizacoes_terceiro_{n}", terceiro),
                    (f"multas_indenizacoes_tipo_{n}", tipo), (f"multas_indenizacoes_vcto_{n}", vencimento),
                    (f"multas_indenizacoes_valor_{n}", valor), (f"multas_indenizacoes_base_pje-calc_{n}", base_multa),
                    (f"multas_indenizacoes_aliquota_{n}", aliquota), (f"multas_indenizacoes_aplicar_juros_{n}", "SIM" if juros else "NÃO"),
                    (f"multas_indenizacoes_a_partir_de_{n}", juros_desde)]
        if indice_outro:
            familia.append((f"multas_indenizacoes_indice_outro_{n}", indice_outro))
        _usar_campos_da_multa(r, n - 1, item, tipo, credor, juros)
    r.familias.append(("multas", familia))


# ---------------------------------------------------------------- bloco novo: advogados

LIMITE_OAB = 9  # o campo OAB da tela do PJe-Calc aceita 9 caracteres
OAB = re.compile(r"^\d{1,6}[A-Z]?/[A-Z]{2}$")
TIPOS_DOCUMENTO_ADVOGADO = {"CPF": 11, "CNPJ": 14, "CEI": 12}


def cpf_valido(digitos):
    if len(digitos) != 11 or digitos == digitos[0] * 11:
        return False
    for tamanho in (9, 10):
        soma = sum(int(d) * (tamanho + 1 - i) for i, d in enumerate(digitos[:tamanho]))
        if (soma * 10 % 11) % 10 != int(digitos[tamanho]):
            return False
    return True


def cnpj_valido(digitos):
    if len(digitos) != 14 or digitos == digitos[0] * 14:
        return False
    for tamanho, pesos in ((12, [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]), (13, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])):
        resto = sum(int(d) * p for d, p in zip(digitos[:tamanho], pesos)) % 11
        if (0 if resto < 2 else 11 - resto) != int(digitos[tamanho]):
            return False
    return True


def pis_valido(digitos):
    if len(digitos) != 11 or len(set(digitos)) == 1:
        return False
    soma = sum(int(d) * p for d, p in zip(digitos[:10], (3, 2, 9, 8, 7, 6, 5, 4, 3, 2)))
    resto = 11 - soma % 11
    return int(digitos[10]) == (0 if resto in (10, 11) else resto)


def _data_liquidacao(r):
    """dataLiquidacao (raiz): o robô liquida nessa data em vez da regra do dia (padrão 44: retificação, esclarecimento etc.
    mantêm a data do cálculo originário). Depois do último dia do mês corrente → 400."""
    if r.dados.get("dataLiquidacao") in (None, ""):
        return
    serial = r.data("dataLiquidacao", r.dados["dataLiquidacao"])
    if serial is None:
        return
    data = ORIGEM_EXCEL + timedelta(days=serial)
    hoje = datetime.now()
    fim_do_mes = (datetime(hoje.year + hoje.month // 12, hoje.month % 12 + 1, 1) - timedelta(days=1))
    if data.date() > fim_do_mes.date():
        r.erro("dataLiquidacao", r.dados["dataLiquidacao"], f"depois do último dia do mês corrente ({fim_do_mes:%d/%m/%Y})")
        return
    r.usar("dataLiquidacao", "data_liquidacao", f"{data:%d/%m/%Y}")


def _parametros_do_calculo(r):
    """Bloco `parametrosDoCalculo` (raiz): exceções da carga horária e do sábado, na aba Parâmetros do Cálculo.

    Vão como texto .json [[início, fim(, valor)], …] em excecoes_carga_horaria / excecoes_sabado. Períodos dentro do
    cálculo (início/fim do cálculo, ou admissão/demissão), fim ≥ início, sem sobreposição dentro de cada lista.
    """
    bloco = r.dados.get("parametrosDoCalculo")
    if bloco is None:
        return
    if not isinstance(bloco, dict):
        r.erro("parametrosDoCalculo", bloco, "deve ser um objeto { excecoesCargaHoraria, excecoesSabado }")
        return
    inicio = r.valores.get("inicio_calculo") or r.valores.get("data_admissao")
    fim = r.valores.get("termino_calculo") or r.valores.get("data_rescisao")
    for nome, chave, com_valor in (("excecoesCargaHoraria", "excecoes_carga_horaria", True),
                                   ("excecoesSabado", "excecoes_sabado", False)):
        caminho = f"parametrosDoCalculo.{nome}"
        lista = bloco.get(nome)
        if lista in (None, []):
            continue
        if not isinstance(lista, list) or not all(isinstance(i, dict) for i in lista):
            r.erro(caminho, lista, "deve ser uma lista de { inicio, fim" + (", valor }" if com_valor else " }"))
            continue
        linhas, erros_antes = [], len(r.erros)
        for n, item in enumerate(lista, start=1):
            onde = f"{caminho}[{n}]"
            extras = set(item) - {"inicio", "fim"} - ({"valor"} if com_valor else set())
            if extras:
                r.erro(onde, item, f"campo(s) não reconhecido(s): {', '.join(sorted(extras))}")
                continue
            a, b = r.data(f"{onde}.inicio", item.get("inicio")), r.data(f"{onde}.fim", item.get("fim"))
            if a is None or b is None:
                continue
            if b < a:
                r.erro(onde, item, "fim anterior ao início")
                continue
            if isinstance(inicio, int) and isinstance(fim, int) and not (inicio <= a and b <= fim):
                r.erro(onde, item, f"fora do período do cálculo ({ORIGEM_EXCEL + timedelta(days=inicio):%d/%m/%Y} a "
                                   f"{ORIGEM_EXCEL + timedelta(days=fim):%d/%m/%Y})")
                continue
            linha = [a, b]
            if com_valor:
                valor = numero_da_planilha(item.get("valor"))
                if valor is None or not 0 < valor <= 744:
                    r.erro(f"{onde}.valor", item.get("valor"), "carga horária mensal deve ser número maior que 0 (ex.: 180)")
                    continue
                linha.append(float(valor))
            linhas.append(linha)
        ordenadas = sorted(linhas)
        for anterior, seguinte in zip(ordenadas, ordenadas[1:]):
            if seguinte[0] <= anterior[1]:
                r.erro(caminho, None, "períodos sobrepostos: " + " e ".join(
                    f"{ORIGEM_EXCEL + timedelta(days=p[0]):%d/%m/%Y}–{ORIGEM_EXCEL + timedelta(days=p[1]):%d/%m/%Y}" for p in (anterior, seguinte)))
        if len(r.erros) != erros_antes:
            continue
        r.valores[chave] = json.dumps([[f"{ORIGEM_EXCEL + timedelta(days=p[0]):%d/%m/%Y}", f"{ORIGEM_EXCEL + timedelta(days=p[1]):%d/%m/%Y}"] + p[2:]
                                       for p in ordenadas])
        for folha in _folhas({nome: lista}, "parametrosDoCalculo"):
            r.usados.add(folha)
    for chave in bloco:
        if chave not in ("excecoesCargaHoraria", "excecoesSabado"):
            r.ignorar(f"parametrosDoCalculo.{chave}", "campo não reconhecido")


def _advogados(r):
    """Bloco `advogados`: { reclamante: [...], reclamado: [...] } → família por parte, lançada em Dados do Cálculo."""
    from servico.verbas_pedido import normalizar_descricao

    bloco = r.dados.get("advogados")
    if bloco is None:
        return
    if not isinstance(bloco, dict):
        r.erro("advogados", bloco, "deve ser um objeto { reclamante: [...], reclamado: [...] }")
        return
    for parte in bloco:
        if parte not in ("reclamante", "reclamado"):
            r.ignorar(f"advogados.{parte}", "parte não reconhecida (use reclamante ou reclamado)")
    for parte in ("reclamante", "reclamado"):
        itens = bloco.get(parte)
        if itens is None:
            continue
        base = f"advogados.{parte}"
        if not isinstance(itens, list):
            r.erro(base, itens, "deve ser uma lista")
            continue
        if len(itens) > 10:
            r.erro(base, len(itens), "no máximo 10 advogados por parte")
            continue
        familia, documentos, erros_antes = [(f"advogados_{parte}_qtd", len(itens))], set(), len(r.erros)
        for n, item in enumerate(itens, start=1):
            caminho = f"{base}[{n - 1}]"
            if not isinstance(item, dict):
                r.erro(caminho, item, "deve ser um objeto { nome, oab, tipoDocumento, documento }")
                continue
            nome = normalizar_descricao(item.get("nome") or "")
            if not nome:
                r.erro(f"{caminho}.nome", item.get("nome"), "obrigatório")
            elif len(nome) > 150:
                r.erro(f"{caminho}.nome", item.get("nome"), "até 150 caracteres (limite da tela)")
            oab = str(item.get("oab") or "").strip().upper().replace(" ", "")
            if oab and (len(oab) > LIMITE_OAB or not OAB.match(oab)):
                r.erro(f"{caminho}.oab", item.get("oab"), f"use NUMERO/UF, até {LIMITE_OAB} caracteres (ex.: 12345/PE, 12345A/PE)")
            elif oab and oab.split("/")[1] not in UFS:
                r.erro(f"{caminho}.oab", item.get("oab"), "UF da OAB não reconhecida")
            tipo = item.get("tipoDocumento")
            documento = documento_da_planilha(item.get("documento")) if item.get("documento") not in (None, "") else None
            if (tipo in (None, "")) != (documento is None):
                r.erro(f"{caminho}.documento", item.get("documento"), "tipoDocumento e documento vêm juntos ou não vêm")
            elif tipo is not None:
                if tipo not in TIPOS_DOCUMENTO_ADVOGADO:
                    r.erro(f"{caminho}.tipoDocumento", tipo, "valor não reconhecido", sorted(TIPOS_DOCUMENTO_ADVOGADO))
                elif (len(documento) != TIPOS_DOCUMENTO_ADVOGADO[tipo]
                      or (tipo == "CPF" and not cpf_valido(documento)) or (tipo == "CNPJ" and not cnpj_valido(documento))):
                    r.erro(f"{caminho}.documento", item.get("documento"), f"{tipo} inválido (tamanho ou dígito verificador)")
                elif documento in documentos:
                    # o PJe-Calc não inclui o segundo advogado com o mesmo documento na mesma parte, e não avisa
                    r.erro(f"{caminho}.documento", item.get("documento"), "documento repetido nesta parte: o PJe-Calc não inclui o segundo")
                documentos.add(documento)
            familia += [(f"advogado_{parte}_{n}_nome", nome), (f"advogado_{parte}_{n}_oab", oab or "-"),
                        (f"advogado_{parte}_{n}_tipodoc", tipo or "-"), (f"advogado_{parte}_{n}_doc", documento or "-")]
        if len(r.erros) == erros_antes and itens:
            r.familias.append((f"advogados_{parte}", familia))
            r.usados.add(base)


# ---------------------------------------------------------------- bloco novo: arquivos (CSV)

CABECALHO_FALTAS = ["INICIO", "FIM", "JUSTIFICADA", "REINICIAR_PER_AQ", "JUSTIFICATIVA"]
CABECALHO_FERIAS = ["RELATIVAS", "PRAZO", "SITUACAO", "DOBRA", "ABONO", "QTD_DIAS_ABONO",
                    "G1INI", "G1FIM", "G1DOBRA", "G2INI", "G2FIM", "G2DOBRA", "G3INI", "G3FIM", "G3DOBRA"]
ARQUIVOS_CONHECIDOS = ("faltas", "ferias", "historico-salarial", "cartao-de-ponto", "base-inss")
DATA_BR = re.compile(r"^\d{2}/\d{2}/\d{4}$")


def _caminho_do_arquivo(r, chave, valor, pasta_do_json):
    caminho = os.path.normpath(str(valor).strip())
    if not os.path.isabs(caminho):
        caminho = os.path.normpath(os.path.join(pasta_do_json, caminho))
    base = f"arquivos.{chave}"
    if Path(caminho).suffix.lower() != ".csv":
        r.erro(base, valor, "deve ser o caminho de um arquivo .csv")
        return None
    if config.RAIZES_PERMITIDAS:
        absoluto = os.path.normcase(os.path.abspath(caminho))
        if not any(absoluto.startswith(os.path.normcase(os.path.abspath(raiz)) + os.sep) for raiz in config.RAIZES_PERMITIDAS):
            r.erro(base, valor, "fora das pastas permitidas para o serviço")
            return None
    if not os.path.isfile(caminho):
        r.erro(base, valor, "arquivo não encontrado")
        return None
    return caminho


def _linhas_do_csv(caminho):
    bruto = Path(caminho).read_bytes()
    try:
        texto = bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = bruto.decode("cp1252")
    return list(csv.reader(texto.splitlines(), delimiter=";"))


def _ler_csv_da_tela(r, chave, caminho, cabecalho, obrigatoria, colunas_data):
    """Faltas e Férias: o .csv já vem no layout das abas PJeFaltas/PJeFérias; devolve as linhas na ordem do cabeçalho."""
    base = f"arquivos.{chave}"
    try:
        linhas = _linhas_do_csv(caminho)
    except OSError as erro:
        r.erro(base, caminho, f"não foi possível ler o arquivo: {erro}")
        return []
    if not linhas:
        r.erro(base, caminho, "arquivo vazio")
        return []
    titulos = [t.strip().lstrip("\ufeff").upper() for t in linhas[0]]
    faltam = [c for c in cabecalho if c not in titulos]
    if faltam:
        r.erro(base, caminho, f"cabeçalho sem as colunas: {', '.join(faltam)}", cabecalho)
        return []
    posicao = {c: titulos.index(c) for c in cabecalho}
    resultado = []
    for numero, linha in enumerate(linhas[1:], start=2):
        if not any(c.strip() for c in linha):
            continue
        celulas = {c: (linha[posicao[c]].strip() if posicao[c] < len(linha) else "") for c in cabecalho}
        if not celulas[obrigatoria]:
            r.erro(base, caminho, f"linha {numero}: '{obrigatoria}' vazio")
            continue
        for coluna in colunas_data:
            if celulas[coluna] and not DATA_BR.match(celulas[coluna]):
                r.erro(base, caminho, f"linha {numero}: '{coluna}' = {celulas[coluna]!r}; use dd/mm/aaaa")
        resultado.append([celulas[c] or None for c in cabecalho])
    return resultado


FLAG_SN = re.compile(r"^[SN]$")
VALOR_BR = re.compile(r"^-?\d+(?:\.\d{3})*(?:,\d+)?$|^-?\d+(?:\.\d+)?$")
COMPETENCIA = re.compile(r"^(0[1-9]|1[0-2])/\d{4}$")
MAXIMO_ERROS_POR_ARQUIVO = 5


def _ler_historico_csv(r, chave, caminho):
    """Histórico salarial no layout da importação do PJe-Calc: MES_ANO;<NOME>;FGTS;FGTS_REC.;CONTRIBUICAO_SOCIAL;CONTRIBUICAO_SOCIAL_REC.

    A 2ª coluna traz o nome do histórico (o SAP manda o próprio nome; o exemplo do PJe-Calc manda 'VALOR').
    Devolve (nome, caminho) ou None se o arquivo tiver erro.
    """
    base = f"arquivos.{chave}"
    try:
        linhas = _linhas_do_csv(caminho)
    except OSError as erro:
        r.erro(base, caminho, f"não foi possível ler o arquivo: {erro}")
        return None
    if not linhas or len(linhas[0]) < 2 or not linhas[0][1].strip():
        r.erro(base, caminho, "cabeçalho inválido: a 1ª coluna é a competência (MM/AAAA) e a 2ª traz o nome do histórico salarial")
        return None
    nome = linhas[0][1].strip().lstrip("\ufeff")[:120]
    if nome.upper() == "VALOR":
        r.erro(base, caminho, "o t\u00edtulo da 2\u00aa coluna \u00e9 o NOME do hist\u00f3rico salarial (ex.: 'ADICIONAL NOTURNO'); "
                              "'VALOR' \u00e9 o t\u00edtulo gen\u00e9rico do exemplo do PJe-Calc")
        return None
    erros_antes, validas = len(r.erros), 0
    for numero, linha in enumerate(linhas[1:], start=2):
        if not any(c.strip() for c in linha):
            continue
        celulas = [c.strip() for c in linha] + [""] * 6
        problema = None
        if not COMPETENCIA.match(celulas[0]):
            problema = f"competência {celulas[0]!r}; use MM/AAAA"
        elif not VALOR_BR.match(celulas[1]):
            problema = f"valor {celulas[1]!r}; use número, com vírgula decimal (1234,56)"
        else:
            ruins = [c for c in celulas[2:6] if c and not FLAG_SN.match(c.upper())]
            if ruins:
                problema = f"marcas FGTS/INSS {ruins!r}; use S ou N"
        if problema:
            if len(r.erros) - erros_antes < MAXIMO_ERROS_POR_ARQUIVO:
                r.erro(base, caminho, f"linha {numero}: {problema}")
            continue
        validas += 1
    if validas == 0 and len(r.erros) == erros_antes:
        r.erro(base, caminho, "arquivo sem nenhuma linha de competência")
    return None if len(r.erros) > erros_antes else (nome, caminho)


HORA = re.compile(r"^([01]\d|2[0-3]):[0-5]\d$")
CABECALHO_CARTAO = ["DATA"] + [f"{t}{n}" for n in range(1, 7) for t in ("ENTRADA", "SAIDA")]


def _ler_cartao_de_ponto(r, chave, caminho):
    """Batidas no layout do ExemploJornadaCartaoDePonto.csv do PJe-Calc (Cartão de Ponto › Grade › Importar Jornada):
    Data;Entrada1;Saída1;…;Entrada6;Saída6 — data dd/mm/aaaa, horas HH:MM. O PJe-Calc recusa o arquivo inteiro se um turno
    tiver só a entrada ou só a saída; a porta recusa antes, citando as datas. Devolve (primeira data, última data) ou None."""
    base = f"arquivos.{chave}"
    try:
        linhas = _linhas_do_csv(caminho)
    except OSError as erro:
        r.erro(base, caminho, f"não foi possível ler o arquivo: {erro}")
        return None
    titulos = [_sem_acento(t.strip().lstrip("\ufeff")).upper().replace(" ", "") for t in (linhas[0] if linhas else [])]
    if titulos[:13] != CABECALHO_CARTAO:
        r.erro(base, caminho, "cabeçalho inválido: use o do exemplo do PJe-Calc (Data;Entrada1;Saída1;…;Entrada6;Saída6)",
               ["Data", "Entrada1", "Saída1", "…", "Entrada6", "Saída6"])
        return None
    erros_antes, problemas, datas = len(r.erros), 0, []
    for numero, linha in enumerate(linhas[1:], start=2):
        if not any(c.strip() for c in linha):
            continue
        celulas = [c.strip() for c in linha] + [""] * 13
        problema = None
        try:
            dia = datetime.strptime(celulas[0], "%d/%m/%Y") if DATA_BR.match(celulas[0]) else None
        except ValueError:
            dia = None
        if dia is None:
            problema = f"data {celulas[0]!r}; use dd/mm/aaaa"
        else:
            for turno in range(6):
                entrada, saida = celulas[1 + 2 * turno], celulas[2 + 2 * turno]
                ruins = [h for h in (entrada, saida) if h and not HORA.match(h)]
                if ruins:
                    problema = f"{celulas[0]}: hora {ruins[0]!r}; use HH:MM"
                    break
                if bool(entrada) != bool(saida):
                    problema = f"{celulas[0]}: turno {turno + 1} com {'entrada sem saída' if entrada else 'saída sem entrada'}"
                    break
        if problema:
            problemas += 1
            if problemas <= MAXIMO_ERROS_POR_ARQUIVO:
                r.erro(base, caminho, f"linha {numero}: {problema}")
            continue
        datas.append(dia)
    if problemas > MAXIMO_ERROS_POR_ARQUIVO:
        r.erro(base, caminho, f"{problemas} linhas com problema no total (mostradas as {MAXIMO_ERROS_POR_ARQUIVO} primeiras); "
                              "o PJe-Calc recusa o arquivo inteiro")
    if not datas and len(r.erros) == erros_antes:
        r.erro(base, caminho, "arquivo sem nenhum dia")
    if len(r.erros) > erros_antes:
        return None
    return min(datas).strftime("%d/%m/%Y"), max(datas).strftime("%d/%m/%Y")


def _valor_br(texto):
    """'1.500,00' / '850,5' / '850' / '1.500' (milhar) -> float."""
    texto = texto.strip()
    if "," in texto:
        return float(texto.replace(".", "").replace(",", "."))
    if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", texto):
        return float(texto.replace(".", ""))
    return float(texto)


def _ler_base_inss(r, chave, caminho):
    """Base do INSS dos salários pagos (coluna "Salários Pagos (Histórico)" da Contribuição Social › Ocorrências).

    MES_ANO;PARCELA;VALOR — PARCELA = SALARIO ou 13. O 13º fica em dezembro ou no mês da demissão/data final
    (é onde a tela repete a competência). Devolve [(MM/AAAA, parcela, valor)].
    """
    base = f"arquivos.{chave}"
    try:
        linhas = _linhas_do_csv(caminho)
    except OSError as erro:
        r.erro(base, caminho, f"não foi possível ler o arquivo: {erro}")
        return []
    titulos = [t.strip().lstrip("\ufeff").upper() for t in (linhas[0] if linhas else [])]
    if titulos[:3] != ["MES_ANO", "PARCELA", "VALOR"]:
        r.erro(base, caminho, "cabeçalho deve ser MES_ANO;PARCELA;VALOR")
        return []
    admissao = r.valores.get("data_admissao")
    final = r.valores.get("termino_calculo") or r.valores.get("data_rescisao")
    d_admissao = ORIGEM_EXCEL + timedelta(days=admissao) if isinstance(admissao, int) else None
    d_final = ORIGEM_EXCEL + timedelta(days=final) if isinstance(final, int) else None
    resultado, vistos, erros_antes = [], set(), len(r.erros)
    for numero, linha in enumerate(linhas[1:], start=2):
        if not any(c.strip() for c in linha):
            continue
        celulas = [c.strip() for c in linha] + [""] * 3
        competencia, parcela, valor = celulas[0], celulas[1].upper(), celulas[2]
        problema = None
        if not COMPETENCIA.match(competencia):
            problema = f"MES_ANO {competencia!r}; use MM/AAAA"
        elif parcela not in ("SALARIO", "13"):
            problema = f"PARCELA {celulas[1]!r}; use SALARIO ou 13"
        elif not VALOR_BR.match(valor):
            problema = f"VALOR {valor!r}; use número com vírgula decimal (1234,56)"
        elif (competencia, parcela) in vistos:
            problema = f"{parcela} de {competencia} repetido"
        else:
            mes, ano = int(competencia[:2]), int(competencia[3:])
            if d_admissao and d_final and not ((d_admissao.year, d_admissao.month) <= (ano, mes) <= (d_final.year, d_final.month)):
                problema = f"{competencia} fora do contrato ({d_admissao:%m/%Y} a {d_final:%m/%Y})"
            elif parcela == "13" and mes != 12 and not (d_final and (ano, mes) == (d_final.year, d_final.month)):
                problema = f"13 em {competencia}: o 13º fica em dezembro ou no mês da demissão/data final"
        if problema:
            if len(r.erros) - erros_antes < MAXIMO_ERROS_POR_ARQUIVO:
                r.erro(base, caminho, f"linha {numero}: {problema}")
            continue
        vistos.add((competencia, parcela))
        resultado.append((competencia, parcela, _valor_br(valor)))
    if not resultado and len(r.erros) == erros_antes:
        r.erro(base, caminho, "arquivo sem nenhuma linha")
    return resultado if len(r.erros) == erros_antes else []


def _conferir_periodos_de_ferias(r, chave, caminho, linhas):
    """O PJe-Calc recusa a importação inteira se um período aquisitivo terminar depois da demissão/data final
    ("As férias definidas no arquivo não podem ser anterior a Data de Admissão e nem posterior a Data de Demissão/Data Final")."""
    admissao = r.valores.get("data_admissao")
    final = r.valores.get("termino_calculo") or r.valores.get("data_rescisao")
    if not isinstance(admissao, int) or not isinstance(final, int):
        return
    d_admissao, d_final = ORIGEM_EXCEL + timedelta(days=admissao), ORIGEM_EXCEL + timedelta(days=final)
    for numero, linha in enumerate(linhas, start=2):
        achado = re.match(r"^(\d{4})/(\d{4})$", str(linha[0] or ""))
        if not achado:
            r.erro(f"arquivos.{chave}", caminho, f"linha {numero}: RELATIVAS {linha[0]!r}; use AAAA/AAAA")
            continue
        ano = int(achado.group(1))
        dia = 28 if (d_admissao.month, d_admissao.day) == (2, 29) else d_admissao.day
        inicio = datetime(ano, d_admissao.month, dia)
        fim = datetime(ano + 1, d_admissao.month, dia) - timedelta(days=1)
        if inicio < d_admissao or fim > d_final:
            r.erro(f"arquivos.{chave}", caminho,
                   f"linha {numero}: período aquisitivo {linha[0]} ({inicio:%d/%m/%Y} a {fim:%d/%m/%Y}) fora do contrato "
                   f"({d_admissao:%d/%m/%Y} a {d_final:%d/%m/%Y}); o PJe-Calc recusa o arquivo inteiro. "
                   "As férias proporcionais o PJe-Calc calcula sozinho: não as mande no .csv")


SITUACOES_DE_FERIAS = {"G": "Gozadas", "GP": "Gozadas Parcialmente", "I": "Indenizadas", "P": "Perdidas"}


def _conferir_situacoes_de_ferias(r, chave, caminho, linhas):
    """Só G, GP, I e P têm significado no PJe-Calc 2.16.0. Outros códigos ou são recusados por ele ou viram GOZADAS:
    "NG" é aceito e gravado como GOZADAS (testado na tela em 13/10/2026) — o oposto de "não gozadas"."""
    posicao = CABECALHO_FERIAS.index("SITUACAO")
    for numero, linha in enumerate(linhas, start=2):
        situacao = str(linha[posicao] or "").strip().upper()
        if situacao and situacao not in SITUACOES_DE_FERIAS:
            r.erro(f"arquivos.{chave}", caminho,
                   f"linha {numero}: SITUACAO {situacao!r}; use G, GP, I ou P (o PJe-Calc grava {situacao!r} como Gozadas). "
                   "Férias não gozadas e já vencidas: I (indenizadas)", sorted(SITUACOES_DE_FERIAS))


def _arquivos(r, caminho_json):
    bloco = r.dados.get("arquivos")
    if not bloco:
        return
    if not isinstance(bloco, dict):
        r.erro("arquivos", bloco, "deve ser um objeto { faltas, ferias, historico-salarial, cartao-de-ponto }")
        return
    pasta = os.path.dirname(os.path.abspath(caminho_json))
    for chave, valor in bloco.items():
        base = f"arquivos.{chave}"
        if chave not in ARQUIVOS_CONHECIDOS:
            r.ignorar(base, "arquivo sem correspondente na automação")
            continue
        if valor in (None, ""):
            continue
        if chave == "historico-salarial" and isinstance(valor, list) and valor and all(isinstance(v, str) for v in valor):
            caminhos = valor
        elif isinstance(valor, str):
            caminhos = [valor]
        else:
            r.erro(base, valor, "deve ser o caminho do arquivo, em texto")
            continue
        if chave != "historico-salarial" and len(caminhos) != 1:
            r.erro(base, valor, "deve ser o caminho de um único arquivo")
            continue
        resolvidos = [_caminho_do_arquivo(r, chave, v, pasta) for v in caminhos]
        if None in resolvidos:
            continue
        caminho = resolvidos[0]
        if chave == "historico-salarial":
            com_cs = []
            for um in resolvidos:
                lido = _ler_historico_csv(r, chave, um)
                if lido and lido[0].upper() in (h[0].upper() for h in r.historicos):
                    r.erro(base, um, f"já há outro histórico salarial chamado {lido[0]!r} neste pedido")
                elif lido:
                    r.historicos.append(lido)
                    if any(len(l) > 4 and l[4].strip().upper() == "S" for l in _linhas_do_csv(um)[1:]):
                        com_cs.append(lido[0])
            if r.historicos and not com_cs:
                r.avisos.append("arquivos.historico-salarial: nenhum histórico com CONTRIBUICAO_SOCIAL = S; o PJe-Calc usa a última "
                                "remuneração como salário pago na Contribuição Social (marque S nas parcelas do salário de contribuição)")
            r.usados.add(base)
        elif chave == "faltas":
            r.faltas = _ler_csv_da_tela(r, chave, caminho, CABECALHO_FALTAS, "INICIO", ("INICIO", "FIM"))
            r.usados.add(base)
        elif chave == "base-inss":
            r.base_inss = _ler_base_inss(r, chave, caminho)
            r.usados.add(base)
        elif chave == "ferias":
            r.ferias = _ler_csv_da_tela(r, chave, caminho, CABECALHO_FERIAS, "RELATIVAS",
                                        ("G1INI", "G1FIM", "G2INI", "G2FIM", "G3INI", "G3FIM"))
            _conferir_periodos_de_ferias(r, chave, caminho, r.ferias)
            _conferir_situacoes_de_ferias(r, chave, caminho, r.ferias)
            r.usados.add(base)
        elif chave == "cartao-de-ponto":
            periodo = _ler_cartao_de_ponto(r, chave, caminho)
            if periodo:
                # a tela Cartão de Ponto importa o arquivo com os critérios padrão no período das batidas
                r.usar(base, "cartao_ponto_arquivo", caminho)
                r.usar(base, "cartao_ponto_inicio", periodo[0])
                r.usar(base, "cartao_ponto_fim", periodo[1])
        else:
            r.arquivos[chave] = caminho
            r.ignorar(base, "arquivo recebido, mas esta importação ainda não é feita pela automação")


# ---------------------------------------------------------------- bloco novo: telas

INCIDENCIAS_MULTA_FGTS = {
    "Sobre devido": "SOBRE_O_TOTAL_DEVIDO", "Sobre saldo e/ou saque": "SOBRE_DEPOSITADO_SACADO",
    "Sobre diferença": "SOBRE_DIFERENCA", "Sobre devido + saldo e/ou saque": "SOBRE_TOTAL_DEVIDO_MAIS_SAQUE_E_OU_SALDO",
    "Sobre devido - saldo e/ou saque": "SOBRE_TOTAL_DEVIDO_MENOS_SAQUE_E_OU_SALDO",
}


def _fgts_completo(r, fgts):
    """Campos da tela FGTS além de destino e multa. Os da multa só existem com a multa marcada (a tela os desabilita),
    e percentual/incidência só com valor Calculado (com Informado a tela troca por um campo de valor)."""
    base = "telas.fgts"
    multa = r.valores.get("multa_fgts") == "Sim"

    def booleano(campo, chave):
        if campo not in fgts:
            return None
        if not isinstance(fgts[campo], bool):
            r.erro(f"{base}.{campo}", fgts[campo], "use true ou false")
            return None
        r.usar(f"{base}.{campo}", chave, _texto(fgts[campo]))
        return fgts[campo]

    def opcao(campo, opcoes):
        valor = fgts.get(campo)
        if valor in (None, ""):
            return None
        alvo = str(valor).replace("\u2212", "-").strip()
        if isinstance(valor, (int, float)) and not isinstance(valor, bool) and float(valor).is_integer():
            alvo = f"{int(valor)}%"  # 8 \u2192 "8%", 40 \u2192 "40%"
        if alvo in opcoes:
            return alvo
        r.erro(f"{base}.{campo}", valor, "valor não reconhecido", list(opcoes))
        return None

    def so_com_multa(campo):
        if campo in fgts and not multa:
            r.erro(f"{base}.{campo}", fgts[campo], "só com fgts_calcular_multa = true (a tela o desabilita sem a multa)")
            return False
        return campo in fgts

    compor = opcao("fgts_compor_principal", ("Sim", "Não"))
    if compor:
        r.usar(f"{base}.fgts_compor_principal", "compor_principal", _texto(compor == "Sim"))
    aliquota = opcao("fgts_aliquota", ("8%", "2%"))
    r.usar(f"{base}.fgts_aliquota", "fgts_aliquota", aliquota)
    booleano("fgts_multa_lc110", "fgts_multa_lc110")
    booleano("fgts_contribuicao_lc110", "fgts_contribuicao_lc110")

    tipo = opcao("fgts_tipo_valor_multa", ("Calculado", "Informado")) if so_com_multa("fgts_tipo_valor_multa") else None
    r.usar(f"{base}.fgts_tipo_valor_multa", "fgts_tipo_valor_multa", tipo)
    for campo in ("fgts_percentual_multa", "fgts_incidencia_multa"):
        if so_com_multa(campo) and tipo == "Informado":
            r.erro(f"{base}.{campo}", fgts[campo], "com valor Informado a tela não tem percentual nem incidência")
    if so_com_multa("fgts_percentual_multa") and tipo != "Informado":
        r.usar(f"{base}.fgts_percentual_multa", "fgts_percentual_multa", opcao("fgts_percentual_multa", ("40%", "20%")))
    if so_com_multa("fgts_incidencia_multa") and tipo != "Informado":
        incidencia = opcao("fgts_incidencia_multa", tuple(INCIDENCIAS_MULTA_FGTS) + tuple(INCIDENCIAS_MULTA_FGTS.values()))
        r.usar(f"{base}.fgts_incidencia_multa", "fgts_incidencia_multa", INCIDENCIAS_MULTA_FGTS.get(incidencia, incidencia))
    if "fgts_valor_multa_informado" in fgts:
        valor = numero_da_planilha(fgts["fgts_valor_multa_informado"])
        if tipo != "Informado":
            r.erro(f"{base}.fgts_valor_multa_informado", fgts["fgts_valor_multa_informado"], "só com fgts_tipo_valor_multa = Informado")
        elif valor is None:
            r.erro(f"{base}.fgts_valor_multa_informado", fgts["fgts_valor_multa_informado"], "valor em reais deve ser número")
        else:
            r.usar(f"{base}.fgts_valor_multa_informado", "fgts_valor_multa_informado", valor)
    elif tipo == "Informado":
        r.erro(f"{base}.fgts_valor_multa_informado", None, "obrigatório com fgts_tipo_valor_multa = Informado")
    if so_com_multa("fgts_multa_467"):
        booleano("fgts_multa_467", "fgts_multa_467")
    if so_com_multa("fgts_excluir_aviso_multa"):
        booleano("fgts_excluir_aviso_multa", "fgts_excluir_aviso_multa")
    elif multa:
        # sem a chave, fica o padrão da tela (marcado); a automação antiga o desmarcava
        r.valores["fgts_excluir_aviso_multa"] = "PADRAO"

    # pensão alimentícia: a caixa "sobre a multa" só aparece com a multa e a pensão sobre o FGTS marcadas
    pensao = booleano("fgts_pensao_alimenticia", "fgts_pensao_alimenticia")
    if "fgts_pensao_alimenticia_multa" in fgts:
        if fgts["fgts_pensao_alimenticia_multa"] is True and not (multa and pensao):
            r.erro(f"{base}.fgts_pensao_alimenticia_multa", True,
                   "só com fgts_calcular_multa = true e fgts_pensao_alimenticia = true (a tela só mostra a caixa assim)")
        elif multa and pensao:
            booleano("fgts_pensao_alimenticia_multa", "fgts_pensao_alimenticia_multa")
        else:
            r.ignorar(f"{base}.fgts_pensao_alimenticia_multa", "false sem multa ou sem pensão: a tela nem mostra a caixa")

    # quadro "Saldo e/ou Saque" (raiz fgtsSaldoSaque) e "Deduzir do FGTS" (uma caixa para o quadro inteiro)
    deduzir = booleano("fgts_deduzir", "fgts_deduzir")
    linhas_bloco = r.dados.get("fgtsSaldoSaque")
    linhas = []
    if linhas_bloco not in (None, []):
        if not isinstance(linhas_bloco, list) or not all(isinstance(i, dict) for i in linhas_bloco):
            r.erro("fgtsSaldoSaque", linhas_bloco, "deve ser uma lista de { data, valor }")
        else:
            admissao = r.valores.get("data_admissao")
            for n, item in enumerate(linhas_bloco):
                onde = f"fgtsSaldoSaque[{n}]"
                if set(item) - {"data", "valor"}:
                    r.erro(onde, item, f"campo(s) não reconhecido(s): {', '.join(sorted(set(item) - {'data', 'valor'}))}")
                    continue
                serial = r.data(f"{onde}.data", item.get("data"))
                valor = numero_da_planilha(item.get("valor"))
                if valor is None or valor <= 0:
                    r.erro(f"{onde}.valor", item.get("valor"), "valor em reais maior que zero")
                elif serial is not None and isinstance(admissao, int) and serial < admissao:
                    r.erro(f"{onde}.data", item.get("data"), "antes da admissão")
                elif serial is not None:
                    linhas.append([f"{ORIGEM_EXCEL + timedelta(days=serial):%d/%m/%Y}", float(valor)])
    incidencia = r.valores.get("fgts_incidencia_multa")
    usa_saldo = multa and incidencia not in (None, "SOBRE_O_TOTAL_DEVIDO")
    if linhas and not deduzir and not usa_saldo:
        r.erro("fgtsSaldoSaque", linhas_bloco, "o PJe-Calc só aceita o quadro como base da multa (incidência com saldo/saque) "
                                               "ou com fgts_deduzir = true")
    elif deduzir and not linhas and not any(e["campo"].startswith("fgtsSaldoSaque") for e in r.erros):
        r.erro(f"{base}.fgts_deduzir", True, "sem linhas em fgtsSaldoSaque a tela deixa a caixa desabilitada")
    elif linhas:
        r.valores["fgts_saldo_saque"] = json.dumps(linhas)
        r.usados.update({"fgtsSaldoSaque[].data", "fgtsSaldoSaque[].valor"})


ALIQUOTAS_SEGURADO = ("Segurado Empregado", "Empregado Doméstico", "Fixa")
ALIQUOTAS_EMPREGADOR = ("Por Atividade Econômica", "Por Período", "Fixa")


def _contribuicao_social_completa(r, cs):
    """telas.contribuicao-social + contribuicaoSocial (raiz). Na tela principal só ficam os quatro checkboxes; o resto
    (alíquotas, períodos de incidência, Simples) é a janela de Ocorrências › Regerar, gravada no Confirmar.

    Os checkboxes e a alíquota do segurado usam as chaves da planilha (cont_social_*); o resto vai em cs_* e as duas
    listas como texto .json. Campo fora da condição em que a tela o mostra → 400.
    """
    base = "telas.contribuicao-social"
    raiz = r.dados.get("contribuicaoSocial")
    raiz = raiz if isinstance(raiz, dict) else {}
    if "contribuicaoSocial" in r.dados and not isinstance(r.dados["contribuicaoSocial"], dict):
        r.erro("contribuicaoSocial", r.dados["contribuicaoSocial"], "deve ser um objeto { simplesNacional, aliquotasPorPeriodo }")

    def booleano(campo, chave, condicao=True, motivo=""):
        if campo not in cs:
            return None
        if not isinstance(cs[campo], bool):
            r.erro(f"{base}.{campo}", cs[campo], "use true ou false")
            return None
        if not condicao:
            r.erro(f"{base}.{campo}", cs[campo], motivo)
            return None
        r.usar(f"{base}.{campo}", chave, _texto(cs[campo]))
        return cs[campo]

    def opcao(campo, opcoes):
        valor = cs.get(campo)
        if valor in (None, ""):
            return None
        if str(valor).strip() in opcoes:
            return str(valor).strip()
        r.erro(f"{base}.{campo}", valor, "valor não reconhecido", list(opcoes))
        return None

    def percentual(caminho, valor):
        numero = numero_da_planilha(valor)
        if numero is None or not 0 <= numero <= 100:
            r.erro(caminho, valor, "percentual deve ser número de 0 a 100")
            return None
        return float(numero)

    def so_com(campo, condicao, motivo):
        if campo in cs and not condicao:
            r.erro(f"{base}.{campo}", cs[campo], motivo)
            return False
        return campo in cs

    # tela principal
    apurar = booleano("cs_apurar_segurado", "cont_social_apurar_segurado")
    sem_apurar = "só com cs_apurar_segurado = true (a tela o desabilita sem ele)"
    cobrar = booleano("cs_cobrar_reclamante", "cont_social_cobrar_reclamante", apurar is not False, sem_apurar)
    booleano("cs_corrigir_desconto_reclamante", "cont_social_com_correcao_trabalhista", apurar is not False, sem_apurar)
    if cobrar is False and cs.get("cs_corrigir_desconto_reclamante") is True:
        # "Com Correção Trabalhista" só aparece na tela com "Cobrar do Reclamante" marcado
        r.erro(f"{base}.cs_corrigir_desconto_reclamante", True, "só com cs_cobrar_reclamante = true (a tela o esconde sem ele)")
    booleano("cs_apurar_salarios_pagos", "cont_social_salario_pagos")

    # janela Regerar: segurado
    segurado = opcao("cs_aliquota_segurado", ALIQUOTAS_SEGURADO)
    r.usar(f"{base}.cs_aliquota_segurado", "cont_social_aliquota_segurado", segurado)
    so_fixa = "só com cs_aliquota_segurado = Fixa"
    if so_com("cs_aliquota_segurado_fixa", segurado == "Fixa", so_fixa):
        r.usar(f"{base}.cs_aliquota_segurado_fixa", "cs_aliquota_segurado_fixa",
               percentual(f"{base}.cs_aliquota_segurado_fixa", cs["cs_aliquota_segurado_fixa"]))
    elif segurado == "Fixa":
        r.erro(f"{base}.cs_aliquota_segurado_fixa", None, "obrigatório com cs_aliquota_segurado = Fixa")
    booleano("cs_limitar_teto", "cs_limitar_teto", segurado == "Fixa", so_fixa)

    # janela Regerar: empregador
    empregador = opcao("cs_aliquota_empregador", ALIQUOTAS_EMPREGADOR)
    r.usar(f"{base}.cs_aliquota_empregador", "cs_aliquota_empregador", empregador)
    por_atividade = "só com cs_aliquota_empregador = Por Atividade Econômica"
    for campo, chave in (("cs_apurar_empresa", "cont_social_apurar_empresa"), ("cs_apurar_sat", "cont_social_apurar_sat"),
                         ("cs_apurar_terceiros", "cont_social_apurar_terceiros")):
        booleano(campo, chave, empregador in (None, "Por Atividade Econômica"), por_atividade)
    if "cs_atividade_economica" in cs and empregador not in (None, "Por Atividade Econômica"):
        r.erro(f"{base}.cs_atividade_economica", cs["cs_atividade_economica"], por_atividade)
    elif empregador == "Por Atividade Econômica" and not r.valores.get("cont_social_aliquota_empregador"):
        r.erro(f"{base}.cs_atividade_economica", None,
               "obrigatório com cs_aliquota_empregador = Por Atividade Econômica (ou dadosDoCalculo.atividadeEconomica)")
    elif empregador not in (None, "Por Atividade Econômica") and r.valores.get("cont_social_aliquota_empregador"):
        r.avisos.append(f"Contribuição Social: alíquota do empregador '{empregador}'; a atividade econômica "
                        f"'{r.valores['cont_social_aliquota_empregador']}' não é usada")
    for campo in ("cs_cota_empresa", "cs_sat", "cs_terceiros"):
        if so_com(campo, empregador == "Fixa", "só com cs_aliquota_empregador = Fixa") and cs[campo] not in (None, ""):
            r.usar(f"{base}.{campo}", campo, percentual(f"{base}.{campo}", cs[campo]))

    # contrato, para as datas e as listas
    admissao = r.valores.get("data_admissao")
    final = r.valores.get("termino_calculo") or r.valores.get("data_rescisao")
    d_admissao = ORIGEM_EXCEL + timedelta(days=admissao) if isinstance(admissao, int) else None
    d_final = ORIGEM_EXCEL + timedelta(days=final) if isinstance(final, int) else None

    for grupo in ("devidos", "pagos"):
        datas = {}
        for ponta in ("inicio", "fim"):
            campo = f"cs_incidencia_{grupo}_{ponta}"
            if cs.get(campo) in (None, ""):
                continue
            serial = r.data(f"{base}.{campo}", cs[campo])
            if serial is not None:
                datas[ponta] = ORIGEM_EXCEL + timedelta(days=serial)
                r.usar(f"{base}.{campo}", campo, f"{datas[ponta]:%d/%m/%Y}")
        if len(datas) == 2 and datas["fim"] < datas["inicio"]:
            r.erro(f"{base}.cs_incidencia_{grupo}_fim", cs[f"cs_incidencia_{grupo}_fim"], "anterior ao início")

    def periodos(nome, extras):
        caminho = f"contribuicaoSocial.{nome}"
        lista = raiz.get(nome)
        if lista in (None, []):
            return []
        if not isinstance(lista, list) or not all(isinstance(i, dict) for i in lista):
            r.erro(caminho, lista, "deve ser uma lista de { inicio, fim, … }")
            return []
        resultado, erros_antes = [], len(r.erros)
        for n, item in enumerate(lista, start=1):
            onde = f"{caminho}[{n}]"
            desconhecidos = set(item) - {"inicio", "fim"} - set(extras)
            if desconhecidos:
                r.erro(onde, item, f"campo(s) não reconhecido(s): {', '.join(sorted(desconhecidos))}")
                continue
            inicio, fim = str(item.get("inicio") or ""), str(item.get("fim") or "")
            if not COMPETENCIA.match(inicio) or not COMPETENCIA.match(fim):
                r.erro(onde, item, "inicio e fim em MM/AAAA")
                continue
            a, b = (int(inicio[3:]), int(inicio[:2])), (int(fim[3:]), int(fim[:2]))
            if b < a:
                r.erro(onde, item, "fim anterior ao início")
                continue
            if d_admissao and d_final and not ((d_admissao.year, d_admissao.month) <= a and b <= (d_final.year, d_final.month)):
                r.erro(onde, item, f"fora do contrato ({d_admissao:%m/%Y} a {d_final:%m/%Y})")
                continue
            valores = []
            for extra, obrigatorio in extras.items():
                if item.get(extra) in (None, ""):
                    if obrigatorio:
                        r.erro(onde, item, f"'{extra}' obrigatório")
                    valores.append(None)
                else:
                    valores.append(percentual(f"{onde}.{extra}", item[extra]))
            resultado.append([inicio, fim, a, b] + valores)
        ordenados = sorted(resultado, key=lambda p: p[2])
        for anterior, seguinte in zip(ordenados, ordenados[1:]):
            if seguinte[2] <= anterior[3]:
                r.erro(caminho, [anterior[:2], seguinte[:2]], f"períodos sobrepostos: {anterior[0]}–{anterior[1]} e {seguinte[0]}–{seguinte[1]}")
        if len(r.erros) != erros_antes:
            return []
        for folha in _folhas({nome: lista}, "contribuicaoSocial"):
            r.usados.add(folha)
        return [[p[0], p[1]] + p[4:] for p in ordenados]

    simples = periodos("simplesNacional", {})
    if simples:
        r.valores["cs_simples_nacional"] = json.dumps(simples)
    por_periodo = periodos("aliquotasPorPeriodo", {"empresa": True, "sat": True, "terceiros": False})
    if raiz.get("aliquotasPorPeriodo") and empregador != "Por Período":
        r.erro("contribuicaoSocial.aliquotasPorPeriodo", raiz["aliquotasPorPeriodo"], "só com cs_aliquota_empregador = Por Período")
    elif por_periodo:
        r.valores["cs_aliquotas_por_periodo"] = json.dumps(por_periodo)
    elif empregador == "Por Período" and not any(e["campo"].startswith("contribuicaoSocial.aliquotasPorPeriodo") for e in r.erros):
        r.erro("contribuicaoSocial.aliquotasPorPeriodo", None, "obrigatório (ao menos um período) com cs_aliquota_empregador = Por Período")
    for chave in raiz:
        if chave not in ("simplesNacional", "aliquotasPorPeriodo"):
            r.ignorar(f"contribuicaoSocial.{chave}", "campo não reconhecido")


INDICES_ACUMULADOS = ("MES_SUBSEQUENTE_AO_VENCIMENTO", "MES_DO_VENCIMENTO", "MES_SUBSEQUENTE_E_MES_DO_VENCIMENTO")


# Correção › Dados Específicos (e "Ignorar Taxa Negativa" de Dados Gerais): padrões da 2.16 num cálculo novo, usados para
# avaliar a condição em que a tela mostra cada campo quando o pedido não manda o campo de que ele depende
PADRAO_DADOS_ESPECIFICOS = {
    "cj_custas_correcao": True,
    "cj_cs_dev_trab_correcao": True, "cj_cs_dev_prev_correcao": False, "cj_cs_dev_multa": False, "cj_cs_lei11941": True,
    "cj_cs_limitar_multa": True,
    "cj_cs_pag_trab_correcao": False, "cj_cs_pag_prev_correcao": True, "cj_cs_pag_multa": True,
}
OPCOES_INDICE_PROPRIO = ("Utilizar índice trabalhista", "Utilizar outro índice")
OPCOES_FGTS = ("Utilizar índice trabalhista", "Utilizar índice JAM", "Utilizar índice JAM e índice trabalhista")
OPCOES_FORMA = ("Mês-a-mês", "A partir de")


def _dados_especificos(r, cj, base):
    """Chaves cj_* da aba Dados Específicos (FGTS, Previdência Privada, Custas, Contribuição Social devidos/pagos) e
    cj_ignorar_taxa_negativa. Vão com o mesmo nome para a planilha; a tela as aplica na ordem em que se habilitam.
    Campo fora da condição em que a tela o mostra → 400 (a condição usa o pedido ou, sem ele, o padrão da 2.16)."""
    listas = CATALOGO["telas"]["correcao-juros-multa"]["listas"]
    efetivo = {}

    def booleano(chave):
        if chave not in cj:
            return PADRAO_DADOS_ESPECIFICOS.get(chave)
        if not isinstance(cj[chave], bool):
            r.erro(f"{base}.{chave}", cj[chave], "use true ou false")
            return PADRAO_DADOS_ESPECIFICOS.get(chave)
        r.usar(f"{base}.{chave}", chave, _texto(cj[chave]))
        return cj[chave]

    def opcao(chave, opcoes):
        if cj.get(chave) in (None, ""):
            return None
        if cj[chave] not in opcoes:
            r.erro(f"{base}.{chave}", cj[chave], "valor não reconhecido", list(opcoes))
            return None
        r.usar(f"{base}.{chave}", chave, cj[chave])
        return cj[chave]

    def data(chave):
        if cj.get(chave) in (None, ""):
            return None
        serial = r.data(f"{base}.{chave}", cj[chave])
        if serial is not None:
            r.usar(f"{base}.{chave}", chave, f"{ORIGEM_EXCEL + timedelta(days=serial):%d/%m/%Y}")
        return serial

    def so_com(chave, condicao, motivo):
        if chave in cj and cj[chave] not in (None, "", False) and not condicao:
            r.erro(f"{base}.{chave}", cj[chave], motivo)
            return False
        return True

    booleano("cj_ignorar_taxa_negativa")
    opcao("cj_fgts_indice", OPCOES_FGTS)
    # Previdência Privada e Custas: índice próprio
    for grupo, condicao in (("pp", True), ("custas", None)):
        if grupo == "custas":
            correcao = booleano("cj_custas_correcao")
            for chave in ("cj_custas_juros", "cj_custas_indice", "cj_custas_outro_indice"):
                so_com(chave, correcao, "só com cj_custas_correcao = true (a tela desabilita sem a correção)")
            booleano("cj_custas_juros")
        else:
            booleano("cj_pp_juros")
        indice = opcao(f"cj_{grupo}_indice", OPCOES_INDICE_PROPRIO)
        outro = cj.get(f"cj_{grupo}_outro_indice")
        if outro not in (None, ""):
            if indice != "Utilizar outro índice":
                r.erro(f"{base}.cj_{grupo}_outro_indice", outro, f"só com cj_{grupo}_indice = \"Utilizar outro índice\"")
            elif outro not in listas["indiceTrabalhista"]:
                r.erro(f"{base}.cj_{grupo}_outro_indice", outro, "fora da lista de índices", listas["indiceTrabalhista"])
            else:
                r.usar(f"{base}.cj_{grupo}_outro_indice", f"cj_{grupo}_outro_indice", outro)
        elif indice == "Utilizar outro índice":
            r.erro(f"{base}.cj_{grupo}_outro_indice", None, f"obrigatório com cj_{grupo}_indice = \"Utilizar outro índice\"")
    # Contribuição Social: salários devidos (dev) e pagos (pag)
    for bloco in ("dev", "pag"):
        p = f"cj_cs_{bloco}_"
        trab, prev = booleano(p + "trab_correcao"), booleano(p + "prev_correcao")
        so_com(p + "trab_juros", trab, f"só com {p}trab_correcao = true")
        booleano(p + "trab_juros")
        so_com(p + "prev_juros", prev, f"só com {p}prev_correcao = true")
        booleano(p + "prev_juros")
        so_com(p + "multa", prev, f"só com {p}prev_correcao = true")
        multa = booleano(p + "multa")
        for chave, opcoes in ((p + "multa_tipo", ("Urbana", "Rural")), (p + "multa_pagamento", ("Integral", "Reduzido"))):
            if so_com(chave, multa and prev, f"só com {p}multa = true"):
                opcao(chave, opcoes)
        forma = None
        if so_com(p + "forma", prev and not trab, f"só com {p}prev_correcao = true e {p}trab_correcao = false"):
            forma = opcao(p + "forma", OPCOES_FORMA)
        if so_com(p + "data", (trab and prev) or forma == "A partir de",
                  f"só com as duas correções ({p}trab_correcao e {p}prev_correcao) ou com {p}forma = \"A partir de\""):
            data(p + "data")
        if forma == "A partir de" and cj.get(p + "data") in (None, ""):
            r.erro(f"{base}.{p}data", None, f"obrigatória com {p}forma = \"A partir de\"")
        # Lei 11.941: a dos devidos usa as chaves de vocês (cj_cs_lei11941*, cj_cs_limitar_multa*); a dos pagos, cj_cs_pag_lei11941*
        lei, lei_data, limitar, limitar_data = (("cj_cs_lei11941", "cj_cs_lei11941_a_partir", "cj_cs_limitar_multa", "cj_cs_limitar_multa_a_partir")
                                                if bloco == "dev" else
                                                ("cj_cs_pag_lei11941", "cj_cs_pag_lei11941_a_partir", "cj_cs_pag_limitar_multa", "cj_cs_pag_limitar_multa_a_partir"))
        visivel = (trab or forma == "A partir de") if bloco == "dev" else forma == "A partir de"
        for chave in (lei, lei_data, limitar, limitar_data):
            so_com(chave, visivel, "a tela só mostra a Lei 11.941 deste bloco " +
                   ("com a correção trabalhista ou com a forma \"A partir de\"" if bloco == "dev" else "com cj_cs_pag_forma = \"A partir de\""))
        if visivel:
            ativa = booleano(lei)
            if so_com(lei_data, ativa, f"só com {lei} = true"):
                data(lei_data)
            if so_com(limitar, ativa, f"só com {lei} = true"):
                limita = booleano(limitar)
                if so_com(limitar_data, limita, f"só com {limitar} = true"):
                    data(limitar_data)


def _correcao_e_juros(r, cj):
    """telas.correcao-juros-multa: tabela de juros, combinações de índices e de juros (várias, para as viradas da
    ADC 58 e da Lei 14.905/2024), base de juros das verbas e índices acumulados (este vai para a tela Liquidar).

    Combinação ausente = padrão do PJe-Calc (com ajuizamento, ele já cria "Sem Juros" a partir dele); lista vazia =
    nenhuma combinação. Valores = texto exato das listas da tela (catalogo_pjecalc.json).
    """
    base = "telas.correcao-juros-multa"
    listas = CATALOGO["telas"]["correcao-juros-multa"]["listas"]
    if cj.get("cj_juros") not in (None, ""):
        if cj["cj_juros"] in listas["juros"]:
            r.usar(f"{base}.cj_juros", "cj_juros", cj["cj_juros"])
        else:
            r.erro(f"{base}.cj_juros", cj["cj_juros"], "tabela de juros fora da lista da tela", listas["juros"])

    def combinacoes(campo, chave_item, opcoes, chave):
        if campo not in cj:
            return
        itens = cj[campo]
        if not isinstance(itens, list) or not all(isinstance(i, dict) for i in itens):
            r.erro(f"{base}.{campo}", itens, f"deve ser uma lista de {{ {chave_item}, a_partir }} (lista vazia = nenhuma combinação)")
            return
        linhas, erros_antes = [], len(r.erros)
        for n, item in enumerate(itens):
            onde = f"{base}.{campo}[{n}]"
            extras = set(item) - {chave_item, "a_partir"}
            if extras:
                r.erro(onde, item, f"campo(s) não reconhecido(s): {', '.join(sorted(extras))}")
                continue
            if item.get(chave_item) not in opcoes:
                r.erro(f"{onde}.{chave_item}", item.get(chave_item), "fora da lista da tela", opcoes)
                continue
            serial = r.data(f"{onde}.a_partir", item.get("a_partir"))
            if serial is not None:
                linhas.append((serial, item[chave_item]))
        datas = [s for s, _ in linhas]
        if len(set(datas)) != len(datas):
            r.erro(f"{base}.{campo}", itens, "duas combinações na mesma data (a tela recusa)")
        if len(r.erros) == erros_antes:
            r.usados.add(f"{base}.{campo}")
            r.valores[chave] = json.dumps([[valor, f"{ORIGEM_EXCEL + timedelta(days=s):%d/%m/%Y}"] for s, valor in sorted(linhas)],
                                          ensure_ascii=False)

    combinacoes("cj_correcoes_combinadas", "indice", listas["indiceTrabalhista"], "cj_correcoes_combinadas")
    combinacoes("cj_juros_combinados", "juros", listas["juros"], "cj_juros_combinados")
    # sem juros pré-judiciais os juros começam no ajuizamento: a tela recusa combinação nessa data ou antes ("O valor do
    # campo 'A partir de' deve ser maior que Data de Ajuizamento") e o cálculo sairia sem juros
    ajuizamento = r.valores.get("data_ajuizamento")
    if cj.get("cj_juros_pre_judicial") is False and isinstance(ajuizamento, int) and r.valores.get("cj_juros_combinados"):
        cedo = [desde for _, desde in json.loads(r.valores["cj_juros_combinados"])
                if (datetime.strptime(desde, "%d/%m/%Y") - ORIGEM_EXCEL).days <= ajuizamento]
        if cedo:
            r.erro(f"{base}.cj_juros_combinados", cj.get("cj_juros_combinados"),
                   f"com cj_juros_pre_judicial = false a combinação precisa ser depois do ajuizamento "
                   f"({ORIGEM_EXCEL + timedelta(days=ajuizamento):%d/%m/%Y}); para juros desde o ajuizamento, use a própria cj_juros")
            r.valores.pop("cj_juros_combinados", None)
    if cj.get("cj_base_juros") not in (None, ""):
        if cj["cj_base_juros"] in listas["baseDeJurosDasVerbas"]:
            r.usar(f"{base}.cj_base_juros", "verbas_base_de_juros", cj["cj_base_juros"])  # chave da planilha (aba Dados Específicos)
        elif _sem_acento(cj["cj_base_juros"]).startswith("padrao"):
            r.ignorar(f"{base}.cj_base_juros", "padrão do sistema: fica o da tela (mande só quando não for o padrão)")
        else:
            r.erro(f"{base}.cj_base_juros", cj["cj_base_juros"], "fora da lista da tela", listas["baseDeJurosDasVerbas"])
    if cj.get("cj_indices_acumulados") not in (None, ""):
        if cj["cj_indices_acumulados"] in INDICES_ACUMULADOS:
            r.usar(f"{base}.cj_indices_acumulados", "cj_indices_acumulados", cj["cj_indices_acumulados"])
        else:
            r.erro(f"{base}.cj_indices_acumulados", cj["cj_indices_acumulados"], "valor não reconhecido", INDICES_ACUMULADOS)
    _dados_especificos(r, cj, base)
    for antigo, motivo in (("cj_outro_indice", "formato antigo: use cj_correcoes_combinadas [{indice, a_partir}]"),
                           ("cj_outro_indice_a_partir", "formato antigo: use cj_correcoes_combinadas [{indice, a_partir}]"),
                           ("cj_fazenda", "fora do contrato: use a opção da lista (juros 'Juros Fazenda Pública' ou índice 'Devedor Fazenda Pública')")):
        if antigo in cj:
            r.ignorar(f"{base}.{antigo}", motivo)


IR_CAIXAS = ("ir_sobre_juros", "ir_cobrar_reclamado", "ir_tributacao_exclusiva", "ir_tributacao_em_separado", "ir_regime_caixa",
             "ir_deduzir_cs_reclamante", "ir_deduzir_previdencia_privada", "ir_deduzir_pensao_alimenticia",
             "ir_deduzir_honorarios_reclamante", "ir_aposentado_65")
IR_FORMATO_ANTIGO = ("ir_regime", "ir_exclusiva", "ir_deducoes")


def _imposto_de_renda(r, ir):
    """telas.imposto-de-renda: as caixas da tela (booleano) e ir_dependentes (inteiro ≥ 0). Com ir_apurar = false a tela
    desabilita todas as outras: mandá-las junto → 400. As chaves do formato antigo (texto) ficam em ignorados."""
    base = "telas.imposto-de-renda"
    apurar = ir.get("ir_apurar")
    if "ir_apurar" in ir:
        if isinstance(apurar, bool):
            r.usar(f"{base}.ir_apurar", "ir_apurar", _texto(apurar))
        else:
            r.erro(f"{base}.ir_apurar", apurar, "use true ou false")
    fora = "só com ir_apurar = true (a tela desabilita as outras caixas sem ele)"
    for campo in IR_CAIXAS:
        if campo not in ir:
            continue
        if not isinstance(ir[campo], bool):
            r.erro(f"{base}.{campo}", ir[campo], "use true ou false")
        elif apurar is False:
            r.erro(f"{base}.{campo}", ir[campo], fora)
        else:
            r.usar(f"{base}.{campo}", campo, _texto(ir[campo]))
    if "ir_dependentes" in ir:
        valor = ir["ir_dependentes"]
        if isinstance(valor, bool) or not isinstance(valor, int) or valor < 0:
            r.erro(f"{base}.ir_dependentes", valor, "número inteiro de 0 em diante")
        elif apurar is False:
            r.erro(f"{base}.ir_dependentes", valor, fora)
        else:
            r.usar(f"{base}.ir_dependentes", "ir_dependentes", valor)
    for campo in IR_FORMATO_ANTIGO:
        if campo in ir:
            r.ignorar(f"{base}.{campo}", "formato antigo (texto); use as caixas do contrato novo (ir_tributacao_exclusiva, …)")


def _telas(r):
    telas = r.dados.get("telas") or {}
    dados = telas.get("dados-do-calculo") or {}
    base = "telas.dados-do-calculo"
    simples = {"reclamante_nome": "nome_parte1", "reclamado_nome": "nome_parte2"}
    for campo, chave in simples.items():
        if dados.get(campo):
            r.usar(f"{base}.{campo}", chave, str(dados[campo])[:150])
    for campo, chave in (("reclamante_cpf", "doc_parte1"), ("reclamado_cnpj", "doc_parte2"), ("reclamado_cpf", "doc_parte2")):
        if dados.get(campo):
            r.usar(f"{base}.{campo}", chave, documento_da_planilha(dados[campo]))
    if "indices_acumulados" in dados:
        r.ignorar(f"{base}.indices_acumulados", "a tela Liquidar tem 3 opções, não sim/não: use telas.correcao-juros-multa."
                                                "cj_indices_acumulados (MES_SUBSEQUENTE_AO_VENCIMENTO, MES_DO_VENCIMENTO, "
                                                "MES_SUBSEQUENTE_E_MES_DO_VENCIMENTO)")
    if dados.get("autuado_em"):
        serial = r.data(f"{base}.autuado_em", dados["autuado_em"])
        if serial is not None:
            r.usar(f"{base}.autuado_em", "autuado_em", f"{ORIGEM_EXCEL + timedelta(days=serial):%d/%m/%Y}")
    if dados.get("reclamante_pis"):
        digitos = re.sub(r"\D", "", str(dados["reclamante_pis"]))
        if pis_valido(digitos):
            # a tela usa a máscara 000.00000.00-0
            r.usar(f"{base}.reclamante_pis", "pis_parte1", f"{digitos[:3]}.{digitos[3:8]}.{digitos[8:10]}-{digitos[10]}")
        else:
            r.erro(f"{base}.reclamante_pis", dados["reclamante_pis"], "PIS inválido (11 dígitos com dígito verificador)")
    if dados.get("numero_processo"):
        if PADRAO_CNJ.match(str(dados["numero_processo"])):
            r.usar(f"{base}.numero_processo", "numero_processo", str(dados["numero_processo"]))
        else:
            r.erro(f"{base}.numero_processo", dados["numero_processo"], "número CNJ fora do formato 0000000-00.0000.0.00.0000")
    for campo, chave in (("admissao", "data_admissao"), ("demissao", "data_rescisao"), ("ajuizamento", "data_ajuizamento"),
                         ("data_inicial", "inicio_calculo"), ("data_final", "termino_calculo")):
        if dados.get(campo):
            r.usar(f"{base}.{campo}", chave, r.data(f"{base}.{campo}", dados[campo]))
    for campo, chave in (("ultima_remuneracao", "ultima_remuneracao"), ("maior_remuneracao", "maior_remuneracao"),
                         ("valor_da_causa", "valor_da_causa")):
        if campo in dados:
            valor = numero_da_planilha(dados[campo])
            if valor is None:
                r.erro(f"{base}.{campo}", dados[campo], "valor em reais deve ser número")
            r.usar(f"{base}.{campo}", chave, valor)
    if "carga_horaria" in dados:
        achado = re.search(r"\d+(?:[.,]\d+)?", str(dados["carga_horaria"]))
        if achado:
            r.usar(f"{base}.carga_horaria", "carga_horaria", float(achado.group(0).replace(",", ".")))
        else:
            r.erro(f"{base}.carga_horaria", dados["carga_horaria"], "carga horária sem número")
    marcacoes = {
        "prescricao_quinquenal": "aplicar_prescricao_verbas", "prescricao_fgts": "aplicar_prescricao_FGTS",
        "projetar_aviso_indenizado": "projetar_aviso_previo_indenizado", "limitar_avos": "limitar_avos_ao_periodo_do_calculo",
        "zerar_valor_negativo": "zerar_valor_negativo_padrao", "feriados_estaduais": "considerar_feriados_estaduais",
        "feriados_municipais": "considerar_feriados_municipais",
        "ponto_facultativo_sexta_feira_santa": "ponto_facultutativo_sextafeirasanta",
        "ponto_facultativo_corpus_christi": "ponto_facultutativo_corpuschristi", "ponto_facultativo_carnaval": "ponto_facultutativo_carnaval",
    }
    for campo, chave in marcacoes.items():
        if campo in dados:
            if isinstance(dados[campo], bool):
                r.usar(f"{base}.{campo}", chave, _texto(dados[campo]))
            else:
                r.erro(f"{base}.{campo}", dados[campo], "use true ou false")
    prazo = None
    if dados.get("prazo_aviso_previo"):
        aceitos = CATALOGO["telas"]["dados-do-calculo"]["listas"]["apuracaoPrazoDoAvisoPrevio"]
        prazo = next((a for a in aceitos if _sem_acento(a) == _sem_acento(dados["prazo_aviso_previo"])), None)
        if prazo is None:
            r.erro(f"{base}.prazo_aviso_previo", dados["prazo_aviso_previo"], "valor não reconhecido", aceitos)
        r.usar(f"{base}.prazo_aviso_previo", "prazo_de_aviso_previo", prazo)
    if "dias_aviso_previo" in dados:
        dias = numero_da_planilha(dados["dias_aviso_previo"])
        if dias is None or dias != int(dias) or dias <= 0:
            r.erro(f"{base}.dias_aviso_previo", dados["dias_aviso_previo"], "número inteiro de dias")
        else:
            r.usar(f"{base}.dias_aviso_previo", "quantidade_dias_aviso_previo", int(dias))
    if prazo == "Informado" and "quantidade_dias_aviso_previo" not in r.valores and not any(
            e["campo"] == f"{base}.dias_aviso_previo" for e in r.erros):
        # O PJe-Calc recusa o formulário ("Campo obrigatório: Quantidade") se o prazo é Informado e não há dias.
        fim_aviso, rescisao = _serial(dados.get("aviso_previo_indenizado")), r.valores.get("data_rescisao")
        if fim_aviso and rescisao and 0 < fim_aviso - rescisao <= 90:
            r.usar(f"{base}.aviso_previo_indenizado", "quantidade_dias_aviso_previo", fim_aviso - rescisao)
            r.avisos.append(f"Dias do aviso prévio deduzidos de aviso_previo_indenizado menos a demissão: {fim_aviso - rescisao}. "
                            "Mande 'dias_aviso_previo' para não depender disso.")
        else:
            r.erro(f"{base}.dias_aviso_previo", dados.get("dias_aviso_previo"),
                   "prazo_aviso_previo 'Informado' exige a quantidade de dias: mande 'dias_aviso_previo' (inteiro)")
    if dados.get("comentarios"):
        if len(str(dados["comentarios"])) > 255:
            r.avisos.append(f"comentarios tem {len(str(dados['comentarios']))} caracteres; o PJe-Calc aceita 255 e o resto foi cortado.")
        r.usar(f"{base}.comentarios", "comentarios_parametros", str(dados["comentarios"])[:255])
    if "sabado_dia_util" in dados:
        r.usar(f"{base}.sabado_dia_util", "sabado_como_dia_util", _texto(bool(dados["sabado_dia_util"])))
    if dados.get("regime") or dados.get("regime_trabalho"):
        campo = "regime" if dados.get("regime") else "regime_trabalho"
        aceitos = CATALOGO["telas"]["dados-do-calculo"]["listas"]["tipoDaBaseTabelada"]
        if dados[campo] in aceitos:
            r.usar(f"{base}.{campo}", "regime_trabalho", dados[campo])
        else:
            r.erro(f"{base}.{campo}", dados[campo], "valor não reconhecido", aceitos)
    if dados.get("uf") and dados.get("estado") and str(dados["uf"]).upper() != str(dados["estado"]).upper():
        r.erro(f"{base}.estado", dados["estado"], f"'estado' diferente de 'uf' ({dados['uf']}); mande só 'uf'")
    elif dados.get("estado"):
        r.usados.add(f"{base}.estado")
    if dados.get("uf") or dados.get("estado") or dados.get("municipio"):
        uf, municipio = str(dados.get("uf") or dados.get("estado") or "").upper(), dados.get("municipio")
        if uf not in UFS:
            r.erro(f"{base}.uf", dados.get("uf"), "UF não reconhecida", sorted(UFS))
        else:
            opcao = opcao_de_municipio(municipio or "", CATALOGO["municipios"].get(uf, []))
            if opcao is None:
                r.erro(f"{base}.municipio", municipio, f"município não está na lista do PJe-Calc para {uf}")
            else:
                r.valores.pop("municipio", None)
                r.usar(f"{base}.uf", "uf", uf)
                r.usar(f"{base}.municipio", "municipio", opcao)

    fgts = telas.get("fgts") or {}
    if "fgts_destino" in fgts:
        texto = _sem_acento(fgts["fgts_destino"])
        destino = "Pagar" if "pagar" in texto else "Recolher" if "recolher" in texto else None
        if destino is None:
            r.erro("telas.fgts.fgts_destino", fgts["fgts_destino"], "valor não reconhecido", ["Pagar", "Recolher à conta vinculada"])
        r.usar("telas.fgts.fgts_destino", "destino_fgts", destino)
    if "fgts_calcular_multa" in fgts:
        r.usar("telas.fgts.fgts_calcular_multa", "multa_fgts", _sim_nao(bool(fgts["fgts_calcular_multa"])))
    _fgts_completo(r, fgts)

    custas = telas.get("custas-judiciais") or {}
    if "custas_reclamado" in custas:
        texto = _sem_acento(custas["custas_reclamado"])
        tipo = ("Não se Aplica" if "nao se aplica" in texto else "Calculada 2%" if "2%" in texto
                else "Informada" if "informad" in texto else None)
        if tipo is None:
            r.erro("telas.custas-judiciais.custas_reclamado", custas["custas_reclamado"], "valor não reconhecido",
                   CATALOGO["telas"]["custas-judiciais"]["opcoes"]["tipoDeCustasDeConhecimentoDoReclamado"])
        elif tipo == "Informada":
            r.erro("telas.custas-judiciais.custas_reclamado", custas["custas_reclamado"], "custas informadas ainda não suportadas pelo .json")
        elif "custas" not in r.dados:
            r.usados.add("telas.custas-judiciais.custas_reclamado")
            r.familias.append(("custas", [("custas_reclamado_conhecimento", tipo), ("custas_recda_conhecimento_info_vcto", "-"),
                                          ("custas_recda_conhecimento_info_valor", "-")]))
        else:
            r.ignorar("telas.custas-judiciais.custas_reclamado", "o bloco 'custas' já define as custas do reclamado")

    contribuicao = telas.get("contribuicao-social") or {}
    if contribuicao.get("cs_atividade_economica"):
        atividade = str(contribuicao["cs_atividade_economica"]).strip()
        anterior = r.valores.get("cont_social_aliquota_empregador")
        if anterior and anterior != atividade:
            r.erro("telas.contribuicao-social.cs_atividade_economica", atividade,
                   f"diferente de dadosDoCalculo.atividadeEconomica ('{anterior}'); mande a mesma atividade nos dois")
        else:
            r.usar("telas.contribuicao-social.cs_atividade_economica", "cont_social_aliquota_empregador", atividade)
    _contribuicao_social_completa(r, contribuicao)

    _imposto_de_renda(r, telas.get("imposto-de-renda") or {})

    correcao = telas.get("correcao-juros-multa") or {}
    if "cj_correcao" in correcao:
        aceitos = CATALOGO["telas"]["correcao-juros-multa"]["listas"]["indiceTrabalhista"]
        valor = str(correcao["cj_correcao"])
        entre_parenteses = re.search(r"\(([^)]+)\)\s*$", valor)
        indice = valor if valor in aceitos else entre_parenteses.group(1) if entre_parenteses and entre_parenteses.group(1) in aceitos else None
        if indice is None:
            r.erro("telas.correcao-juros-multa.cj_correcao", valor, "índice não reconhecido", aceitos)
        r.usar("telas.correcao-juros-multa.cj_correcao", "indice_trabalhista", indice)
    if "cj_juros_pre_judicial" in correcao:
        if isinstance(correcao["cj_juros_pre_judicial"], bool):
            r.usar("telas.correcao-juros-multa.cj_juros_pre_judicial", "aplicar_juros_fase_pre_judicial",
                   _texto(correcao["cj_juros_pre_judicial"]))
        else:
            r.erro("telas.correcao-juros-multa.cj_juros_pre_judicial", correcao["cj_juros_pre_judicial"], "use true ou false")
    _correcao_e_juros(r, correcao)

    for tela, campos in telas.items():
        for campo in (campos or {}):
            caminho = f"telas.{tela}.{campo}"
            if caminho not in r.usados and not any(e["campo"] == caminho for e in r.erros):
                if tela in ("salario-familia", "seguro-desemprego", "previdencia-privada", "pensao-alimenticia",
                            "cartao-de-ponto"):
                    r.ignorar(caminho, "tela ainda não automatizada")
                else:
                    r.ignorar(caminho, "campo sem correspondente na automação")


# Chave que diz se a planilha já traz a família (na v3.41 os blocos vazios têm '-' nela).
PORTADORES = {
    "honorarios": re.compile(r"^(honorarios_periciais_descricao|honorarios_contratuais_tipob)_\d+$"),
    "multas": re.compile(r"^multas_indenizacoes_credor_\d+$"),
    "custas": re.compile(r"^(custas_reclamado_conhecimento|custas_recolhidas_rcdo\d+_valor)$"),
    # a planilha da casa traz um advogado por parte (advogado_parte1/2): se ela traz, prevalece a parte inteira
    "advogados_reclamante": re.compile(r"^advogado_parte1$"),
    "advogados_reclamado": re.compile(r"^advogado_parte2$"),
}


def combinar_com_planilha(interpretacao, existentes):
    """Planilha + .json: devolve (complemento, forçados, avisos).

    Campo avulso: a planilha manda; o .json só preenche o que ela deixou vazio.
    Família (honorários, multas, custas): é um bloco inteiro. Se a planilha já a traz, prevalece inteira; se não,
    a do .json substitui o bloco todo — inclusive os '0' e o 'qtd' da v3.41, que senão sairiam misturados.
    """
    from Calculo.leitura_planilha import valor_vazio

    avisos = []
    complemento = interpretacao.pares_avulsos()
    for chave, valor in complemento:
        atual = existentes.get(chave)
        if chave in existentes and not valor_vazio(atual) and not valor_vazio(valor) and str(atual) != str(valor):
            avisos.append(f"a planilha prevaleceu em '{chave}' (planilha={atual!r}; .json={valor!r})")
    forcados = []
    for nome, familia in interpretacao.familias:
        if any(PORTADORES[nome].match(chave) and not valor_vazio(valor) for chave, valor in existentes.items()):
            avisos.append(f"a planilha já traz {nome}; os do .json foram ignorados")
        else:
            forcados += familia
    return complemento, forcados, avisos


def criar_planilha_do_pedido(caminho_json, pares, execucao_id, faltas=None, ferias=None, historicos=None, base_inss=None):
    """Pedido sem planilha: monta, ao lado do .json, uma planilha mínima que os fluxos já sabem ler.

    PJE-BD com as chaves traduzidas; Faltas e Férias com as linhas dos .csv (ou vazias, 'Excluir Linha'); e a aba SEM_HISTORICO, que faz o
    Histórico Salarial e a base do INSS serem pulados com aviso. O PDF e o .PJC saem na mesma pasta (a do .json).
    """
    import openpyxl
    from Calculo.leitura_planilha import ABA_SEM_HISTORICO, ABA_CSV_HISTORICO, ABA_BASE_INSS

    destino = Path(caminho_json).with_name(f"{execucao_id}_automacao_pedido.xlsx")
    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "PJE-BD"
    aba.append([])
    aba.append(["IDENTIFICADOR", "INFORMACAO"])
    for chave, valor in pares:
        aba.append([chave, valor])
    livro.create_sheet("PJeFaltas").append(["INICIO", "FIM", "JUSTIFICADA", "REINICIAR_PER_AQ", "JUSTIFICATIVA"])
    for linha in faltas or [["Excluir Linha"]]:
        livro["PJeFaltas"].append(linha)
    livro.create_sheet("PJeFérias").append(["RELATIVAS", "PRAZO", "SITUACAO", "DOBRA", "ABONO", "QTD_DIAS_ABONO",
                                            "G1INI", "G1FIM", "G1DOBRA", "G2INI", "G2FIM", "G2DOBRA", "G3INI", "G3FIM", "G3DOBRA"])
    for linha in ferias or [["Excluir Linha"]]:
        livro["PJeFérias"].append(linha)
    livro.create_sheet(ABA_SEM_HISTORICO).append(["Pedido montado só com o .json de lançamentos, sem histórico salarial."])
    if historicos:
        # a tela do Histórico Salarial importa estes .csv (nome do histórico, caminho) no lugar da grade da planilha
        aba_csv = livro.create_sheet(ABA_CSV_HISTORICO)
        for nome, arquivo in historicos:
            aba_csv.append([nome, arquivo])
    if base_inss:
        # a Contribuição Social grava estes valores na coluna "Salários Pagos (Histórico)", casando por competência
        aba_inss = livro.create_sheet(ABA_BASE_INSS)
        for competencia, parcela, valor in base_inss:
            aba_inss.append([competencia, parcela, valor])
    livro.save(destino)
    return str(destino)


def interpretar(caminho):
    """Lê e traduz o .json. Levanta LancamentosIlegiveis se o arquivo não for do contrato; erros de campo ficam em .erros."""
    try:
        dados = json.loads(Path(caminho).read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as erro:
        raise LancamentosIlegiveis(f"não foi possível ler o .json de lançamentos: {erro}")
    if not isinstance(dados, dict):
        raise LancamentosIlegiveis("o .json de lançamentos precisa ser um objeto { … }")
    contrato = str(dados.get("contrato") or "")
    if contrato != CONTRATO and not contrato.startswith(CONTRATO + "."):
        raise LancamentosIlegiveis(f"contrato '{contrato}' não suportado; esperado '{CONTRATO}'")

    r = Interpretacao(dados)
    _dados_do_calculo(r)
    _fgts(r)
    _custas(r)
    _honorarios(r)
    _multas(r)
    _telas(r)
    _parametros_do_calculo(r)
    _data_liquidacao(r)
    _advogados(r)
    _arquivos(r, caminho)
    from servico.criterios_cartao import _cartao_de_ponto
    _cartao_de_ponto(r)  # depois de _arquivos: os critérios só valem com o arquivo do cartão
    if "cont_social_aliquota_empregador" in r.valores:
        # A CNAE só é lida quando a leitura chega em entidade_filantropica; sem valor, nenhuma das três opções é marcada.
        for gatilho in ("processo_de_agroindustria", "empresa_enquadrada_simples", "empresa_com_desoneracao", "entidade_filantropica"):
            r.valores.setdefault(gatilho, "<nao_preenchido>")
    from servico.verbas_pedido import interpretar_verbas
    interpretar_verbas(r)
    from servico.seguro_previdencia import interpretar_seguro_e_previdencia
    interpretar_seguro_e_previdencia(r)  # depois das verbas e dos históricos, que ele referencia
    # ignorados por item (honorarios[0].dataVencimento) já contam como tratados
    tratados = {re.sub(r"\[\d+\]", "[]", k) for k in r.ignorados}
    for folha in _folhas({k: v for k, v in dados.items() if k not in ("telas", "verbas", "arquivos", "reflexosManuais", "advogados", "contribuicaoSocial", "parametrosDoCalculo", "fgtsSaldoSaque", "cartaoDePonto", "seguroDesemprego", "previdenciaPrivada")}):
        if folha and folha not in META and folha not in r.usados and folha not in tratados and not folha.endswith("contrato"):
            r.ignorar(folha, "campo não reconhecido")
    return r
