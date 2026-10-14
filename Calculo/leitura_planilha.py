# -*- coding: utf-8 -*-
"""Leitura tolerante das planilhas: a v3.41 (.xlsb) e a gerada pelo SAP (.xlsm) guardam os mesmos dados em tipos diferentes.

A v3.41 entrega datas como serial do Excel e flags como texto ('True', 'FALSE');
a do SAP entrega datas, booleanos e números de verdade, e usa nomes de aba com espaço ('PJe Faltas').
"""
import os
import re
import unicodedata
from datetime import datetime
from numbers import Real

import pandas as pd
import xlrd

MARCADORES_VAZIO = {"", "-", "<nao_preenchido>", "<não_preenchido>", "<em branco>", "<oculto>"}
TEXTOS_VERDADEIRO = {"true", "verdadeiro", "sim", "s", "1"}
TEXTOS_FALSO = {"false", "falso", "não", "nao", "n", "0"}


def valor_vazio(valor):
    if valor is None or valor is pd.NaT:
        return True
    if isinstance(valor, float) and pd.isna(valor):
        return True
    return isinstance(valor, str) and valor.strip().lower() in MARCADORES_VAZIO


def data_da_planilha(valor):
    if valor_vazio(valor) or isinstance(valor, bool):
        return None
    if isinstance(valor, datetime):
        return pd.Timestamp(valor).to_pydatetime()
    if isinstance(valor, Real):
        return xlrd.xldate_as_datetime(valor, 0)
    if isinstance(valor, str):
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(valor.strip(), formato)
            except ValueError:
                continue
    return None


def booleano_da_planilha(valor):
    if isinstance(valor, bool):
        return valor
    if valor_vazio(valor):
        return None
    texto = str(valor).strip().lower()
    if texto in TEXTOS_VERDADEIRO:
        return True
    if texto in TEXTOS_FALSO:
        return False
    return None


def numero_da_planilha(valor):
    if isinstance(valor, bool) or valor_vazio(valor):
        return None
    if isinstance(valor, Real):
        return float(valor)
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def documento_da_planilha(valor):
    """CPF/CNPJ só com dígitos: aceita número (inclusive float inteiro do .xlsm) ou texto com máscara."""
    if isinstance(valor, bool) or valor_vazio(valor):
        return None
    if isinstance(valor, Real):
        valor = f"{int(valor)}"
    return re.sub(r"\D", "", str(valor)) or None


UFS = {"AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS", "MG", "PA", "PB", "PR", "PE", "PI",
       "RJ", "RN", "RS", "RO", "RR", "SC", "SP", "SE", "TO"}


def _nome_normalizado(texto):
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ASCII", "ignore").decode("ASCII")
    return " ".join(re.sub(r"[^\w\s]", " ", texto).upper().split())


def opcao_de_municipio(valor, opcoes):
    """Acha na lista do PJe-Calc a opção do município: 'RECIFE' (v3.41) ou 'RECIFE-PE' / 'Recife/PE' (SAP)."""
    por_nome = {}
    for opcao in opcoes:
        por_nome.setdefault(_nome_normalizado(opcao), opcao)
    alvo = _nome_normalizado(valor)
    candidatos = [alvo]
    partes = alvo.split()
    if len(partes) > 1 and partes[-1] in UFS:
        candidatos.append(" ".join(partes[:-1]))
    for candidato in candidatos:
        if candidato in por_nome:
            return por_nome[candidato]
    sem_espaco = {nome.replace(" ", ""): opcao for nome, opcao in por_nome.items()}
    for candidato in candidatos:
        if candidato.replace(" ", "") in sem_espaco:
            return sem_espaco[candidato.replace(" ", "")]
    return None


# Grupos em que a escolha entre candidatos MUDA a alíquota RAT (TBATIVIDADEECONOMICA + TBALIQUOTARATPORATIVIDADE,
# PJe-Calc 2.16.0, 1.509 atividades; levantado pelo SAP). Nos outros 96 grupos ambíguos o RAT é o mesmo.
# Em "Impressão de jornais" e "Cultivo de junco" as duas entradas têm o MESMO texto: só o id as distingue.
CNAES_RAT_DIVERGENTE = {
    "impressão de jornais": "3 × 2",
    "clubes sociais, esportivos e similares (647-futebol profissional)": "0 × 2, cooperativa",
    "cultivo de junco": "3 × 2",
}


def _normalizar_atividade(texto):
    return " ".join(str(texto).lower().split())


def _destaque_rat(candidatos):
    for _, texto in candidatos:
        normalizado = _normalizar_atividade(texto)
        for atividade, rats in CNAES_RAT_DIVERGENTE.items():
            if atividade in normalizado:
                return f" ⚠ Neste CNAE as opções têm RAT diferente ({rats}); a escolha altera o cálculo."
    return ""


def escolher_atividade_economica(procurada, opcoes, ambiguo_escolhe_primeiro=False):
    """Escolhe a atividade econômica (CNAE) entre as sugestões do PJe-Calc.

    Vale a igualdade exata; sem ela, o "contém". Havendo mais de um candidato (inclusive entradas de texto idêntico,
    que só o id distingue), não há como escolher pelo texto: levanta ValueError, ou, com ambiguo_escolhe_primeiro
    (janela antiga), fica com o primeiro como antes e avisa no log. Devolve (posição, texto) ou None se nenhuma servir.
    """
    alvo = _normalizar_atividade(procurada)
    candidatos = [(i, texto) for i, texto in enumerate(opcoes) if _normalizar_atividade(texto) == alvo]
    if not candidatos:
        candidatos = [(i, texto) for i, texto in enumerate(opcoes) if alvo in _normalizar_atividade(texto)]
    if not candidatos:
        return None
    if len(candidatos) == 1:
        return candidatos[0]

    lista = "; ".join(texto for _, texto in candidatos)
    destaque = _destaque_rat(candidatos)
    if ambiguo_escolhe_primeiro:
        print(f"- [AVISO][CNAE]: {len(candidatos)} candidatos para '{procurada}' ({lista}); escolhido o primeiro. "
              f"Confira a alíquota RAT; na planilha, prefira a descrição completa.{destaque}")
        return candidatos[0]
    raise ValueError(f"{len(candidatos)} candidatos para a atividade econômica '{procurada}' ({lista}). "
                     f"Informe a descrição completa em cont_social_aliquota_empregador.{destaque}")


def aba_existente(caminho, *nomes):
    """Devolve o primeiro nome de aba que existe na planilha (ex.: 'PJeFaltas' na v3.41, 'PJe Faltas' na do SAP)."""
    with pd.ExcelFile(caminho) as planilha:
        abas = set(planilha.sheet_names)
    for nome in nomes:
        if nome in abas:
            return nome
    raise ValueError(f"A planilha não tem nenhuma das abas: {', '.join(nomes)}")


_GRADES = {}
_COMPLEMENTOS_PJE_BD = {}
ABA_SEM_HISTORICO = "SEM_HISTORICO"
ABA_CSV_HISTORICO = "CSV_HISTORICO"


ABA_BASE_INSS = "BASE_INSS"


def base_inss_do_pedido(caminho):
    """[(MM/AAAA, 'SALARIO'|'13', valor)] da base do INSS enviada no .json (arquivos["base-inss"]); [] se não houver."""
    with pd.ExcelFile(caminho) as planilha:
        if ABA_BASE_INSS not in planilha.sheet_names:
            return []
        tabela = planilha.parse(ABA_BASE_INSS, header=None, dtype=str)
    return [(str(c).strip(), str(p).strip(), float(v)) for c, p, v in zip(tabela[0], tabela[1], tabela[2])
            if not pd.isna(c) and not pd.isna(v)]


def csvs_de_historico(caminho):
    """[(nome, caminho do .csv)] de um pedido montado do .json com arquivos de histórico salarial; [] se não houver."""
    with pd.ExcelFile(caminho) as planilha:
        if ABA_CSV_HISTORICO not in planilha.sheet_names:
            return []
        tabela = planilha.parse(ABA_CSV_HISTORICO, header=None)
    return [(str(nome).strip(), str(arquivo).strip()) for nome, arquivo in zip(tabela[0], tabela[1])
            if not pd.isna(nome) and not pd.isna(arquivo)]


class SemHistoricoSalarial(ValueError):
    """Pedido montado só com o .json de lançamentos, sem histórico salarial: as telas que dependem dele são puladas."""


def definir_complemento_pje_bd(caminho, pares, forcados=()):
    """Registra (chave, valor) do .json de lançamentos para esta planilha: 'pares' só preenchem o que está vazio;
    'forcados' substituem o valor da planilha (famílias inteiras que a planilha não traz)."""
    _COMPLEMENTOS_PJE_BD[os.path.abspath(caminho)] = [(c, v, False) for c, v in pares] + [(c, v, True) for c, v in forcados]


def limpar_complementos():
    _COMPLEMENTOS_PJE_BD.clear()


def ler_pje_bd(caminho):
    """A PJE-BD da planilha, completada pelo .json de lançamentos: a planilha manda no que ela carrega.

    Chave existente com valor vazio recebe o do .json no mesmo lugar (a ordem das linhas importa: há chaves de disparo);
    chave que a planilha não tem entra no fim, na ordem em que veio.
    """
    tabela = pd.read_excel(caminho, sheet_name="PJE-BD", header=1)
    pares = _COMPLEMENTOS_PJE_BD.get(os.path.abspath(caminho))
    if not pares:
        return tabela
    tabela = tabela.astype({"INFORMACAO": object})
    posicao = {str(chave).strip(): i for i, chave in enumerate(tabela["IDENTIFICADOR"]) if isinstance(chave, str)}
    novas = []
    for chave, valor, forcar in pares:
        if chave in posicao:
            atual = tabela.at[tabela.index[posicao[chave]], "INFORMACAO"]
            if forcar or valor_vazio(atual):
                tabela.at[tabela.index[posicao[chave]], "INFORMACAO"] = valor
        else:
            novas.append({"IDENTIFICADOR": chave, "INFORMACAO": valor})
    if novas:
        tabela = pd.concat([tabela, pd.DataFrame(novas)], ignore_index=True)
    return tabela.reset_index(drop=True)


ABA_HIST_VAL = "PJE HIST-VAL"
ABA_PJE_HISTORICO = "PJe Histórico"
LARGURA_HIST_VAL = 47
ORIGEM_EXCEL = datetime(1899, 12, 30)


def _serial_excel(data):
    return (data - ORIGEM_EXCEL).days


def _competencias_da_coluna(celula, linha_inicial, coluna):
    """Lê datas de competência para baixo até a primeira célula que não é data."""
    linha, datas = linha_inicial, []
    while isinstance(celula(linha, coluna), datetime):
        datas.append((linha, celula(linha, coluna)))
        linha += 1
    return datas


def _grade_hist_val_da_pje_historico(caminho):
    """Monta, com as posições da PJE HIST-VAL da v3.41, os dados da aba PJe Histórico gerada pelo SAP.

    PJe Histórico (linhas achadas pelos rótulos): 'Nº de ordem' em B com a ordem em C:V; abaixo, o nome e o incide FGTS
    (SIM/NÃO); depois uma competência por linha em B, com os valores em C:V. Em Y:AA, sob 'Competência'/'Parcela'/
    'Base INSS', a competência, a parcela (Salário/13º) e a base do INSS.
    PJE HIST-VAL: D5:W5 nome, D6:W6 'SIM' ou 0, C7↓ competência (serial) e D:W valores, C5 quantidade;
    AA5 quantidade do INSS, AA7↓ competência e AB7↓ base; AH4/AK4/AN4 = 0 (sem históricos de depósito de FGTS).
    """
    import openpyxl

    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        aba = livro[ABA_PJE_HISTORICO]
        linhas = [list(l) for l in aba.iter_rows(min_row=1, max_row=aba.max_row, min_col=1, max_col=27, values_only=True)]
        teto_inss = None
        if "INSS" in livro.sheetnames:
            teto_inss = next(livro["INSS"].iter_rows(min_row=12, max_row=12, min_col=3, max_col=3, values_only=True))[0]
            # INSS!C12 vem da função VBA INSSTeto: zerada, vazia ou com erro (#NAME?) = planilha salva sem rodar as macros.
            if not (isinstance(teto_inss, Real) and not isinstance(teto_inss, bool) and teto_inss > 0):
                raise ValueError(f"A planilha foi salva sem calcular as macros (INSS!C12 = {teto_inss!r}): "
                                 "abra no Excel com as macros habilitadas, confira e salve antes de enviar.")
    finally:
        livro.close()

    def celula(linha, coluna):
        if linha - 1 < len(linhas) and coluna - 1 < len(linhas[linha - 1]):
            return linhas[linha - 1][coluna - 1]
        return None

    def texto(linha, coluna):
        return str(celula(linha, coluna) or "").strip()

    # As linhas são achadas pelos rótulos: o gerador do SAP já mudou a altura do cabeçalho (linha 9 → 3).
    ordem = next((l for l in range(1, 41) if texto(l, 2) == "Nº de ordem"), None)
    inss = next((l for l in range(1, 41) if texto(l, 25) == "Competência" and texto(l, 27) == "Base INSS"), None)
    if ordem is None or inss is None or texto(ordem + 1, 2) != "Histórico":
        faltando = [nome for nome, achado in (("'Nº de ordem' na coluna B", ordem), ("'Competência'/'Base INSS' em Y/AA", inss)) if achado is None]
        if ordem is not None and texto(ordem + 1, 2) != "Histórico":
            faltando.append(f"'Histórico' na coluna B, linha {ordem + 1}")
        raise ValueError(f"A aba '{ABA_PJE_HISTORICO}' não está no layout esperado (não achei {', '.join(faltando)}).")
    linha_nome, linha_fgts, primeira = ordem + 1, ordem + 2, ordem + 3

    competencias = _competencias_da_coluna(celula, primeira, 2)
    if not competencias:
        raise ValueError(f"A aba '{ABA_PJE_HISTORICO}' está sem competências calculadas (primeira competência vazia ou zero): "
                         "abra a planilha no Excel, confira e salve antes de enviar.")
    competencias_inss = _competencias_da_coluna(celula, inss + 1, 25)

    historicos = []
    for j in range(20):
        nome = celula(linha_nome, 3 + j)
        nome = str(nome).strip() if nome not in (None, "") else None
        incide_fgts = str(celula(linha_fgts, 3 + j) or "").strip().upper() == "SIM"
        valores = [numero_da_planilha(celula(linha, 3 + j)) or 0.0 for linha, _ in competencias]
        historicos.append((nome, incide_fgts, valores))

    qtd = len(competencias)
    grade = [[None] * LARGURA_HIST_VAL for _ in range(6 + max(qtd, len(competencias_inss)) + 1)]

    grade[3][1], grade[3][2] = "REG", "MÊS/ANO"
    grade[4][2], grade[5][2] = qtd, "BASE FGTS"
    for j, (nome, incide_fgts, valores) in enumerate(historicos):
        grade[3][3 + j] = j + 1
        grade[4][3 + j] = nome
        grade[5][3 + j] = "SIM" if incide_fgts else 0
        for i, valor in enumerate(valores):
            grade[6 + i][3 + j] = float(valor)
    for i, (_, competencia) in enumerate(competencias):
        grade[6 + i][1], grade[6 + i][2] = i + 1, _serial_excel(competencia)

    grade[2][32] = qtd
    for coluna in (34, 35, 37, 38, 40, 41):
        grade[2][coluna] = 0
    for coluna in (33, 36, 39):
        grade[3][coluna] = 0
    grade[4][31], grade[5][31], grade[5][32] = _serial_excel(competencias[0][1]), _serial_excel(competencias[-1][1]), qtd

    grade[4][26] = len(competencias_inss)
    for k, (linha, competencia) in enumerate(competencias_inss):
        grade[6 + k][25], grade[6 + k][26] = k + 1, _serial_excel(competencia)
        grade[6 + k][27] = float(numero_da_planilha(celula(linha, 27)) or 0.0)
    if competencias_inss and not any(grade[6 + k][27] for k in range(len(competencias_inss))):
        print(f"- [AVISO][PJe Histórico]: base do INSS zerada em todas as {len(competencias_inss)} competências "
              f"(teto INSS!C12 = {teto_inss}). Se o processo tem contribuição social, etiquete as rubricas de INSS "
              "na Ficha Financeira.")
    return grade


def _nomes_de_coluna(cabecalho):
    """Reproduz os nomes que o pd.read_excel daria: vazio vira 'Unnamed: n' e repetido ganha '.1', '.2'…"""
    nomes, vistos = [], {}
    for posicao, valor in enumerate(cabecalho):
        nome = f"Unnamed: {posicao}" if valor is None else valor
        if nome in vistos:
            vistos[nome] += 1
            nome = f"{nome}.{vistos[nome]}"
        else:
            vistos[nome] = 0
        nomes.append(nome)
    return nomes


def historicos_da_planilha_sap(caminho):
    """Nomes dos históricos da aba 'PJe Histórico' (planilha gerada pelo SAP); None se a planilha não a tem (v3.41).

    Levanta ValueError se a aba estiver fora do layout ou salva sem calcular as macros.
    """
    with pd.ExcelFile(caminho) as planilha:
        if ABA_PJE_HISTORICO not in planilha.sheet_names:
            return None
    chave = (os.path.abspath(caminho), os.path.getmtime(caminho))
    if chave not in _GRADES:
        _GRADES.clear()
        _GRADES[chave] = _grade_hist_val_da_pje_historico(caminho)
    return [nome for nome in _GRADES[chave][4][3:23] if isinstance(nome, str) and nome.strip()]


def limpar_cache():
    """Chamado no início de cada execução, para os avisos da leitura aparecerem no log de cada uma."""
    _GRADES.clear()


def ler_pje_hist_val(caminho, header):
    """Com a aba PJE HIST-VAL (v3.41), faz exatamente o pd.read_excel de antes; sem ela, usa a PJe Histórico do SAP."""
    with pd.ExcelFile(caminho) as planilha:
        abas = set(planilha.sheet_names)
        if ABA_HIST_VAL in abas:
            return pd.read_excel(planilha, sheet_name=ABA_HIST_VAL, header=header)
    if ABA_PJE_HISTORICO not in abas:
        if ABA_SEM_HISTORICO in abas:
            raise SemHistoricoSalarial("O pedido não trouxe histórico salarial.")
        raise ValueError(f"A planilha não tem a aba '{ABA_HIST_VAL}' nem a '{ABA_PJE_HISTORICO}'.")
    chave = (os.path.abspath(caminho), os.path.getmtime(caminho))
    if chave not in _GRADES:
        _GRADES.clear()
        _GRADES[chave] = _grade_hist_val_da_pje_historico(caminho)
    grade = _GRADES[chave]
    tabela = pd.DataFrame(grade[header + 1:], columns=_nomes_de_coluna(grade[header]))
    return tabela.astype(object).where(tabela.notna(), float("nan"))


def datas_como_texto(tabela, colunas):
    """Converte datas de verdade em 'dd/mm/aaaa'; seriais do Excel e textos ficam como estão."""
    def converter(valor):
        if isinstance(valor, datetime) and not pd.isna(valor):
            return valor.strftime('%d/%m/%Y')
        return valor

    for coluna in colunas:
        if coluna in tabela.columns:
            tabela[coluna] = tabela[coluna].map(converter)
    return tabela
